from __future__ import annotations

from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.services.maintenance_history_import_service import (
    HISTORY_TEMPLATE_HEADERS,
    build_maintenance_history_template,
    parse_maintenance_history_workbook,
)
from app.schemas.inspection_schema import InspectionCreateRequest


def test_downloadable_template_has_controlled_headers_guidance_and_dropdowns() -> None:
    content = build_maintenance_history_template("WP CAA-1010")
    workbook = load_workbook(BytesIO(content))

    assert workbook.sheetnames == ["Maintenance History", "Instructions", "Allowed Values"]
    sheet = workbook["Maintenance History"]
    assert [cell.value for cell in sheet[1]] == list(HISTORY_TEMPLATE_HEADERS)
    assert sheet[2][0].value == "WP CAA-1010"
    assert sheet.freeze_panes == "A2"
    assert len(sheet.data_validations.dataValidation) >= 2
    assert workbook["Instructions"]["A1"].value == "FleetOpt Maintenance History Import"


def test_parser_accepts_completed_history_rows_and_normalizes_values() -> None:
    content = build_maintenance_history_template("WP CAA-1010")
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["Maintenance History"]
    sheet.append([
        "WP CAA-1010", "2026-01-15", 12000, "braking", "scheduled_service",
        "Brake pads inspected and replaced", "A. Perera", "Colombo Workshop", 28500, "Routine service",
    ])
    payload = BytesIO()
    workbook.save(payload)

    rows = parse_maintenance_history_workbook(payload.getvalue(), "WP CAA-1010", current_mileage=15000)

    assert len(rows) == 1
    assert rows[0]["maintenance_system"] == "braking"
    assert rows[0]["event_kind"] == "scheduled_service"
    assert rows[0]["status"] == "completed"
    assert rows[0]["record_source"] == "imported_history"


def test_parser_rejects_entire_workbook_with_row_level_errors() -> None:
    content = build_maintenance_history_template("WP CAA-1010")
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["Maintenance History"]
    sheet.append(["WRONG-REG", "not-a-date", 20000, "unknown", "bad-kind", "", "", "", -10, ""])
    payload = BytesIO()
    workbook.save(payload)

    with pytest.raises(HTTPException) as exc:
        parse_maintenance_history_workbook(payload.getvalue(), "WP CAA-1010", current_mileage=15000)

    assert exc.value.status_code == 422
    detail = exc.value.detail
    assert detail["message"] == "Maintenance history workbook contains invalid rows"
    assert detail["errors"][0]["row"] == 3


def test_commissioning_inspection_does_not_require_an_assessment() -> None:
    payload = InspectionCreateRequest(
        vehicle_id="64b64b64b64b64b64b64b64b",
        inspection_type="commissioning",
        priority="medium",
        inspection_areas=["General condition", "Braking", "Tires"],
    )

    assert payload.assessment_id is None
    assert payload.feature_hash is None
