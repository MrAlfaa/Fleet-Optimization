from datetime import datetime, timezone
from io import BytesIO

import pytest
from bson import ObjectId
from openpyxl import load_workbook

from app.models.delivery_model import DELIVERY_COLLECTION
from app.models.demand_model import (
    DEMAND_FORECAST_COLLECTION,
    FLEET_RECOMMENDATION_COLLECTION,
    OPERATIONAL_DEMAND_DAILY_COLLECTION,
)
from app.models.route_model import ROUTE_PLAN_COLLECTION, ROUTE_SEGMENT_COLLECTION
from app.models.user_model import USER_COLLECTION
from app.models.vehicle_model import VEHICLE_COLLECTION
from app.services import report_service


NOW = datetime(2026, 8, 28, 8, tzinfo=timezone.utc)


def _rows_by_collection():
    run_id, route_id, vehicle_id, driver_id, order_id = (ObjectId() for _ in range(5))
    return run_id, {
        DEMAND_FORECAST_COLLECTION: [
            {"_id": ObjectId(), "model_run_id": run_id, "zone_id": "colombo_central", "forecast_date": NOW,
             "predicted_demand": 42, "lower_bound": 36, "upper_bound": 49},
            {"_id": ObjectId(), "model_run_id": run_id, "zone_id": "borella", "forecast_date": NOW,
             "predicted_demand": 18, "lower_bound": 14, "upper_bound": 23},
        ],
        FLEET_RECOMMENDATION_COLLECTION: [
            {"_id": ObjectId(), "model_run_id": run_id, "zone_id": "colombo_central", "forecast_date": NOW,
             "predicted_demand": 42, "demand_level": "high", "estimated_required_vehicles": 5,
             "recommended_action": "Stage additional delivery capacity"},
            {"_id": ObjectId(), "model_run_id": run_id, "zone_id": "borella", "forecast_date": NOW,
             "predicted_demand": 18, "demand_level": "medium", "estimated_required_vehicles": 2,
             "recommended_action": "Use planned capacity"},
        ],
        OPERATIONAL_DEMAND_DAILY_COLLECTION: [
            {"_id": ObjectId(), "operational_zone_id": "colombo_central", "demand_date": NOW, "total_orders": 40},
        ],
        ROUTE_PLAN_COLLECTION: [
            {"_id": route_id, "route_code": "ROUTE-001", "vehicle_id": vehicle_id, "driver_id": driver_id,
             "delivery_date": "2026-08-28", "delivery_period": "first_half", "delivery_order_ids": [order_id],
             "optimized_stop_order": [{"sequence": 1, "type": "delivery", "order_id": str(order_id),
                                        "point_index": 1, "order_code": "ORD-001", "zone_id": "colombo_central"}],
             "total_predicted_eta_minutes": 38.5, "average_delay_probability": 0.22,
             "optimization_method": "Predictive ETA OR-Tools single vehicle", "status": "dispatched",
             "created_at": NOW, "dispatched_at": NOW},
        ],
        ROUTE_SEGMENT_COLLECTION: [
            {"_id": ObjectId(), "route_plan_id": route_id, "to_stop_index": 1, "distance_km": 8.4,
             "predicted_eta_minutes": 18.2, "delay_probability": 0.22, "delay_risk_level": "low", "created_at": NOW},
        ],
        DELIVERY_COLLECTION: [
            {"_id": order_id, "order_code": "ORD-001", "zone_id": "colombo_central", "priority": "high",
             "parcel_weight": 12.5, "delivery_status": "delivered", "created_at": NOW},
        ],
        VEHICLE_COLLECTION: [
            {"_id": vehicle_id, "registration_number": "WP CAA-1208", "vehicle_code": "QA-EV-008"},
        ],
        USER_COLLECTION: [{"_id": driver_id, "full_name": "Nimal Perera"}],
    }


def _matrix(payload: bytes, sheet_name: str):
    workbook = load_workbook(BytesIO(payload), data_only=True)
    return [[cell.value for cell in row] for row in workbook[sheet_name].iter_rows()]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("report_type", "sheet", "headers"),
    [
        ("demand_forecast", "Demand Forecast", {"Forecast Date", "Zone", "Forecast Orders", "Recorded Orders", "Variance", "Required Vehicle Shifts"}),
        ("demand_capacity", "Capacity Summary", {"Zone", "Forecast Days", "Total Forecast Orders", "Peak Demand", "Maximum Vehicle Shifts"}),
        ("route_optimization", "Route Optimization", {"Route Code", "Delivery Date", "Vehicle", "Driver", "Total Distance (km)", "Average Delay Risk %"}),
        ("route_stop_performance", "Route Stops", {"Route Code", "Stop Sequence", "Order Code", "Delivery Status", "Leg ETA (minutes)", "Delay Risk %"}),
    ],
)
async def test_planning_reports_export_business_ready_workbooks(monkeypatch, report_type, sheet, headers):
    run_id, rows = _rows_by_collection()

    async def fake_find(collection, query, fields, **kwargs):
        return rows.get(collection, [])

    monkeypatch.setattr(report_service, "_find_rows", fake_find)
    monkeypatch.setattr(report_service, "_active_demand_run_id", lambda: _async_value(run_id))

    payload = await report_service.build_specialized_xlsx_report(report_type=report_type)
    matrix = _matrix(payload, sheet)

    assert headers <= set(matrix[4])
    assert len(matrix) >= 6
    assert "model_version" not in str(matrix).lower()
    assert "synthetic" not in str(matrix).lower()


@pytest.mark.asyncio
async def test_planning_report_filters_apply_to_demand_and_routes(monkeypatch):
    run_id, rows = _rows_by_collection()

    async def fake_find(collection, query, fields, **kwargs):
        return rows.get(collection, [])

    monkeypatch.setattr(report_service, "_find_rows", fake_find)
    monkeypatch.setattr(report_service, "_active_demand_run_id", lambda: _async_value(run_id))

    demand = await report_service.build_specialized_xlsx_report(
        report_type="demand_forecast", zone_id="colombo_central", demand_level="high"
    )
    routes = await report_service.build_specialized_xlsx_report(
        report_type="route_optimization", route_status="completed", delivery_period="second_half"
    )

    assert len(_matrix(demand, "Demand Forecast")) == 6
    assert len(_matrix(routes, "Route Optimization")) == 5


async def _async_value(value):
    return value
