from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from bson import ObjectId
from fastapi import HTTPException
from openpyxl import load_workbook

from app.models.inspection_model import INSPECTION_COLLECTION
from app.models.maintenance_model import MAINTENANCE_COLLECTION
from app.models.predictive_model import PREDICTION_COLLECTION
from app.models.user_model import USER_COLLECTION
from app.models.vehicle_model import VEHICLE_COLLECTION
from app.routes import report_routes
from app.services import report_service


NOW = datetime(2026, 8, 27, 10, 30, tzinfo=timezone.utc)


def _matches(row: dict, query: dict) -> bool:
    for key, expected in query.items():
        actual = row.get(key)
        if isinstance(expected, dict) and "$in" in expected:
            if actual not in expected["$in"]:
                return False
        elif actual != expected:
            return False
    return True


class _Cursor:
    def __init__(self, rows: list[dict]):
        self.rows = rows

    def sort(self, field: str, direction: int):
        self.rows.sort(key=lambda row: row.get(field) or datetime.min.replace(tzinfo=timezone.utc), reverse=direction < 0)
        return self

    def limit(self, count: int):
        self.rows = self.rows[:count]
        return self

    async def to_list(self, length: int):
        return deepcopy(self.rows[:length])


class _Collection:
    def __init__(self, rows: list[dict] | None = None):
        self.rows = deepcopy(rows or [])

    def find(self, query: dict, projection: dict | None = None):
        rows = [deepcopy(row) for row in self.rows if _matches(row, query)]
        if projection:
            included = {key for key, enabled in projection.items() if enabled}
            rows = [
                {key: value for key, value in row.items() if key in included or key == "_id"}
                for row in rows
            ]
        return _Cursor(rows)


def _database() -> tuple[dict, ObjectId, ObjectId]:
    vehicle_id = ObjectId()
    second_vehicle_id = ObjectId()
    assignee_id = ObjectId()
    inspection_id = ObjectId()
    assessment_id = ObjectId()
    db = {
        VEHICLE_COLLECTION: _Collection([
            {
                "_id": vehicle_id,
                "registration_number": "WP CAA-1208",
                "vehicle_code": "QA-EV-008",
                "vehicle_category": "bus",
                "fuel_type": "electric",
                "brand": "=EV Corp",
                "model": "K9",
                "vehicle_type": "electric_bus",
                "year": 2023,
                "in_service_date": date(2024, 1, 15),
                "payload_capacity_kg": 4200,
                "current_mileage": 68400,
                "status": "under_maintenance",
                "health_status": "warning",
                "assigned_driver_id": ObjectId(),
                "created_at": datetime(2025, 1, 10, tzinfo=timezone.utc),
                "updated_at": NOW,
            },
            {
                "_id": second_vehicle_id,
                "registration_number": "WP CAB-4401",
                "vehicle_code": "QA-VAN-004",
                "vehicle_category": "van",
                "fuel_type": "diesel",
                "brand": "Safe Motors",
                "model": "Cargo",
                "vehicle_type": "diesel_van",
                "year": 2021,
                "current_mileage": 112000,
                "status": "active",
                "health_status": "healthy",
                "created_at": datetime(2024, 2, 1, tzinfo=timezone.utc),
                "updated_at": NOW,
            },
        ]),
        PREDICTION_COLLECTION: _Collection([
            {
                "_id": assessment_id,
                "vehicle_id": vehicle_id,
                "risk_level": "critical",
                "failure_probability_30_days": 0.27,
                "failure_probability_2000_km": 0.12,
                "system_risks": [
                    {"maintenance_system": "battery_system"},
                    {"maintenance_system": "braking"},
                ],
                "planning_status": "planning_window_available",
                "planning_window": {
                    "days": {"lower": 12, "upper": 30},
                    "km": {"lower": 900, "upper": 2200},
                },
                "data_quality": "good",
                "dispatch_hold": True,
                "feature_hash": "must-not-export",
                "model_version": "must-not-export",
                "created_at": datetime(2026, 8, 26, tzinfo=timezone.utc),
            },
            {
                "_id": ObjectId(),
                "vehicle_id": second_vehicle_id,
                "risk_level": "low",
                "failure_probability_30_days": 0.01,
                "failure_probability_2000_km": 0.01,
                "system_risks": [],
                "planning_status": "no_near_term_failure_indicated",
                "planning_window": None,
                "data_quality": "fair",
                "dispatch_hold": False,
                "created_at": datetime(2026, 7, 1, tzinfo=timezone.utc),
            },
        ]),
        INSPECTION_COLLECTION: _Collection([
            {
                "_id": inspection_id,
                "vehicle_id": vehicle_id,
                "assessment_id": assessment_id,
                "status": "maintenance_required",
                "priority": "critical",
                "inspection_areas": ["battery_system", "braking"],
                "assigned_to": assignee_id,
                "started_at": datetime(2026, 8, 26, 8, 0, tzinfo=timezone.utc),
                "completed_at": datetime(2026, 8, 26, 9, 30, tzinfo=timezone.utc),
                "checklist": [
                    {"outcome": "passed", "remark": None},
                    {"outcome": "failed", "remark": "=Replace module"},
                    {"outcome": "not_applicable", "remark": None},
                ],
                "summary": "+Battery service required",
                "decision": "maintenance_required",
                "decision_notes": "Repair before assignment",
                "created_at": datetime(2026, 8, 26, 7, 45, tzinfo=timezone.utc),
                "updated_at": NOW,
            },
            {
                "_id": ObjectId(),
                "vehicle_id": second_vehicle_id,
                "status": "returned_to_service",
                "priority": "low",
                "inspection_areas": ["tires"],
                "checklist": [{"outcome": "passed", "remark": None}],
                "summary": "No issue found",
                "decision": "returned_to_service",
                "created_at": datetime(2026, 7, 2, tzinfo=timezone.utc),
                "completed_at": datetime(2026, 7, 2, 1, tzinfo=timezone.utc),
            },
        ]),
        MAINTENANCE_COLLECTION: _Collection([
            {
                "_id": ObjectId(),
                "vehicle_id": vehicle_id,
                "inspection_id": inspection_id,
                "maintenance_system": "battery_system",
                "work_order_status": "in_progress",
                "status": "in_progress",
                "priority": "critical",
                "service_date": datetime(2026, 8, 27, tzinfo=timezone.utc),
                "created_at": datetime(2026, 8, 26, 10, tzinfo=timezone.utc),
                "started_at": datetime(2026, 8, 27, 8, tzinfo=timezone.utc),
            }
        ]),
        USER_COLLECTION: _Collection([
            {"_id": assignee_id, "full_name": "Nimal Perera", "email": "hidden@example.com", "password_hash": "hidden"}
        ]),
    }
    return db, vehicle_id, second_vehicle_id


def _sheet_matrix(payload: bytes, sheet_name: str) -> list[list[object]]:
    workbook = load_workbook(BytesIO(payload), data_only=True)
    return [[cell.value for cell in row] for row in workbook[sheet_name].iter_rows()]


def _all_workbook_text(payload: bytes) -> str:
    workbook = load_workbook(BytesIO(payload), data_only=True)
    return "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("report_type", "sheet_name", "required_headers"),
    [
        ("vehicle_maintenance", "Vehicle Maintenance", {"Registration No", "Assessment Result", "30-Day Service Risk %", "Recommended Next Action"}),
        ("vehicle_details", "Vehicle Details", {"Registration No", "Category", "Current Mileage", "Created Date"}),
        ("daily_inspection", "Daily Inspection", {"Inspection Created Date", "Assigned Inspector", "Failed Checks", "Decision"}),
    ],
)
async def test_specialized_reports_have_business_headers_and_safe_content(
    monkeypatch, report_type: str, sheet_name: str, required_headers: set[str]
) -> None:
    db, _, _ = _database()
    monkeypatch.setattr(report_service, "_db", lambda: db)

    payload = await report_service.build_specialized_xlsx_report(report_type=report_type)
    matrix = _sheet_matrix(payload, sheet_name)
    headers = set(matrix[4])
    text = _all_workbook_text(payload)

    assert required_headers <= headers
    if report_type == "vehicle_details":
        assert "'=EV Corp" in text
    elif report_type == "daily_inspection":
        assert "'+Battery service required" in text
    for prohibited in ("feature_hash", "model_version", "password_hash", "synthetic", "prediction"):
        assert prohibited not in text.lower()


@pytest.mark.asyncio
async def test_vehicle_maintenance_filters_vehicle_category_risk_workflow_and_inspection(monkeypatch) -> None:
    db, vehicle_id, _ = _database()
    monkeypatch.setattr(report_service, "_db", lambda: db)

    payload = await report_service.build_specialized_xlsx_report(
        report_type="vehicle_maintenance",
        start_date=datetime(2026, 8, 25, tzinfo=timezone.utc),
        end_date=datetime(2026, 8, 28, tzinfo=timezone.utc),
        vehicle_id=str(vehicle_id),
        vehicle_category="bus",
        assessment_result="critical",
        workflow_status="maintenance",
        inspection_status="maintenance_required",
    )
    matrix = _sheet_matrix(payload, "Vehicle Maintenance")
    headers, rows = matrix[4], matrix[5:]
    registration_index = headers.index("Registration No")
    result_index = headers.index("Assessment Result")
    workbook = load_workbook(BytesIO(payload), data_only=True)
    sheet = workbook["Vehicle Maintenance"]
    risk_column = headers.index("30-Day Service Risk %") + 1

    assert [row[registration_index] for row in rows] == ["WP CAA-1208"]
    assert rows[0][result_index] == "Immediate Service Review"
    assert "vehicle_id=" in str(matrix[2][1])
    assert "Assessment created, inspection created/completed, or service activity date" in str(matrix[2][1])
    assert sheet.freeze_panes == "A6"
    assert sheet.auto_filter.ref.endswith(str(sheet.max_row))
    assert sheet.cell(6, risk_column).number_format == "0.00"


@pytest.mark.asyncio
async def test_vehicle_details_category_and_created_date_filters(monkeypatch) -> None:
    db, _, second_vehicle_id = _database()
    monkeypatch.setattr(report_service, "_db", lambda: db)

    payload = await report_service.build_specialized_xlsx_report(
        report_type="vehicle_details",
        start_date=datetime(2024, 1, 1, tzinfo=timezone.utc),
        end_date=datetime(2024, 12, 31, 23, 59, tzinfo=timezone.utc),
        vehicle_id=str(second_vehicle_id),
        vehicle_category="van",
    )
    matrix = _sheet_matrix(payload, "Vehicle Details")

    assert len(matrix[5:]) == 1
    assert "Vehicle created date" in str(matrix[2][1])
    assert "WP CAB-4401" in _all_workbook_text(payload)


@pytest.mark.asyncio
async def test_inspection_reports_filter_status_and_build_summary_counts(monkeypatch) -> None:
    db, vehicle_id, _ = _database()
    monkeypatch.setattr(report_service, "_db", lambda: db)

    daily = await report_service.build_specialized_xlsx_report(
        report_type="daily_inspection",
        vehicle_id=str(vehicle_id),
        inspection_status="maintenance_required",
        start_date=datetime(2026, 8, 26, tzinfo=timezone.utc),
        end_date=datetime(2026, 8, 26, 23, 59, tzinfo=timezone.utc),
    )
    summary = await report_service.build_specialized_xlsx_report(report_type="inspection_summary")

    daily_matrix = _sheet_matrix(daily, "Daily Inspection")
    assert len(daily_matrix[5:]) == 1
    assert "Nimal Perera" in _all_workbook_text(daily)

    workbook = load_workbook(BytesIO(summary), data_only=True)
    assert workbook.sheetnames == ["Summary", "Inspection Details"]
    summary_text = _all_workbook_text(summary)
    for status in ("Pending", "Assigned", "In Progress", "Decision Pending", "Maintenance Required", "Returned To Service", "Cancelled", "Completed"):
        assert status in summary_text


@pytest.mark.asyncio
async def test_vehicle_maintenance_uses_current_inspection_over_inactive_historical_hold(monkeypatch) -> None:
    db, vehicle_id, _ = _database()
    historical = db[INSPECTION_COLLECTION].rows[0]
    historical["is_active"] = False
    db[INSPECTION_COLLECTION].rows.append({
        "_id": ObjectId(),
        "vehicle_id": vehicle_id,
        "status": "returned_to_service",
        "is_active": False,
        "priority": "low",
        "inspection_areas": ["tires"],
        "created_at": datetime(2026, 8, 27, 12, tzinfo=timezone.utc),
        "completed_at": datetime(2026, 8, 27, 13, tzinfo=timezone.utc),
    })
    db[MAINTENANCE_COLLECTION].rows[0]["work_order_status"] = "completed"
    db[MAINTENANCE_COLLECTION].rows[0]["status"] = "completed"
    monkeypatch.setattr(report_service, "_db", lambda: db)

    payload = await report_service.build_specialized_xlsx_report(
        report_type="vehicle_maintenance", vehicle_id=str(vehicle_id)
    )
    matrix = _sheet_matrix(payload, "Vehicle Maintenance")
    headers, row = matrix[4], matrix[5]

    assert row[headers.index("Inspection Status")] == "Returned To Service"


@pytest.mark.asyncio
async def test_report_route_wires_specialized_query_and_deterministic_filename(monkeypatch) -> None:
    captured: dict = {}

    async def fake_builder(**kwargs):
        captured.update(kwargs)
        return b"xlsx"

    monkeypatch.setattr(report_routes, "build_specialized_xlsx_report", fake_builder)
    response = await report_routes.export_report(
        format="xlsx",
        report_type="vehicle_details",
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 31),
        vehicle_id=None,
        vehicle_category="bus",
        assessment_result=None,
        workflow_status=None,
        inspection_status=None,
        _={"role": "admin"},
    )

    assert captured["report_type"] == "vehicle_details"
    assert captured["vehicle_category"] == "bus"
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["content-disposition"] == "attachment; filename=fleetopt_vehicle_details_2026-08-01_to_2026-08-31.xlsx"


@pytest.mark.asyncio
async def test_invalid_specialized_report_combinations_return_422() -> None:
    with pytest.raises(HTTPException) as pdf_error:
        await report_routes.export_report(
            format="pdf", report_type="vehicle_maintenance", start_date=None, end_date=None,
            vehicle_id=None, vehicle_category=None, assessment_result=None,
            workflow_status=None, inspection_status=None, _={"role": "admin"},
        )
    assert pdf_error.value.status_code == 422

    with pytest.raises(HTTPException) as vehicle_error:
        await report_service.build_specialized_xlsx_report(
            report_type="vehicle_details", vehicle_id="not-an-object-id"
        )
    assert vehicle_error.value.status_code == 422


def test_reports_route_keeps_permission_dependency() -> None:
    route = next(route for route in report_routes.router.routes if route.path == "/reports/export")
    assert route.dependant.dependencies
    assert report_routes.REPORT_ROLES == ["admin", "fleet_manager", "maintenance_manager", "route_planner"]
