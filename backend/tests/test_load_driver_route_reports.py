from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

import pytest
from bson import ObjectId
from openpyxl import load_workbook

from app.models.delivery_model import DELIVERY_COLLECTION
from app.models.driver_device_model import DRIVER_DEVICE_COLLECTION
from app.models.load_optimization_model import PARCEL_CONSOLIDATIONS_COLLECTION
from app.models.route_model import ROUTE_PLAN_COLLECTION, ROUTE_SEGMENT_COLLECTION
from app.models.user_model import USER_COLLECTION
from app.models.vehicle_model import VEHICLE_COLLECTION
from app.services import report_service


NOW = datetime(2026, 8, 29, 8, 30, tzinfo=timezone.utc)


def _matrix(payload: bytes, sheet: str) -> list[list[object]]:
    workbook = load_workbook(BytesIO(payload), data_only=True)
    return [[cell.value for cell in row] for row in workbook[sheet].iter_rows()]


@pytest.mark.asyncio
async def test_consolidation_report_exports_one_filterable_row_per_parcel(monkeypatch) -> None:
    parcel_id = ObjectId()
    rows = {
        PARCEL_CONSOLIDATIONS_COLLECTION: [{
            "_id": ObjectId(), "consolidation_id": "PCID202608291H001", "delivery_date": "2026-08-29",
            "delivery_period": "first_half", "status": "routed", "parcel_count": 1,
            "parcel_ids": [str(parcel_id)], "parcel_points": [{"parcel_id": str(parcel_id), "region_id": 3}],
            "active_plan_id": "LOADPLAN-001", "active_route_set_id": "RSET-001", "created_at": NOW,
            "updated_at": NOW,
        }],
        DELIVERY_COLLECTION: [{
            "_id": parcel_id, "order_code": "ORD-001", "zone_id": "colombo_central",
            "parcel_weight": 14.5, "delivery_status": "assigned", "load_planning_status": "routed",
            "load_profile": {"fragile": True},
        }],
    }

    async def fake_find(collection, query, fields, **kwargs):
        return rows.get(collection, [])

    monkeypatch.setattr(report_service, "_find_rows", fake_find)
    payload = await report_service.build_specialized_xlsx_report(
        report_type="consolidation_details", consolidation_id="PCID202608291H001",
        consolidation_status="routed", delivery_period="first_half",
    )
    matrix = _matrix(payload, "Consolidations")
    headers = matrix[4]
    row = matrix[5]

    assert {"Consolidation ID", "Current Status", "Order Code", "Region", "Load Plan ID", "Route Set ID"} <= set(headers)
    assert row[headers.index("Consolidation ID")] == "PCID202608291H001"
    assert row[headers.index("Order Code")] == "ORD-001"
    assert row[headers.index("Fragile")] == "Yes"


@pytest.mark.asyncio
async def test_driver_details_report_excludes_credentials_and_summarizes_route_activity(monkeypatch) -> None:
    driver_id, vehicle_id, route_id = ObjectId(), ObjectId(), ObjectId()
    rows = {
        USER_COLLECTION: [{
            "_id": driver_id, "role": "driver", "full_name": "Nimal Perera", "username": "nimal",
            "email": "nimal@example.com", "phone": "0770000000", "is_active": True,
            "device_enrollment_status": "enrolled", "password_hash": "never-export", "created_at": NOW,
        }],
        DRIVER_DEVICE_COLLECTION: [{"_id": ObjectId(), "user_id": driver_id, "status": "active", "device_model": "Pixel", "last_seen_at": NOW}],
        ROUTE_PLAN_COLLECTION: [{
            "_id": route_id, "driver_id": driver_id, "vehicle_id": vehicle_id, "route_code": "ROUTE-001",
            "status": "completed", "delivery_order_ids": [ObjectId(), ObjectId()], "completed_at": NOW,
        }],
        VEHICLE_COLLECTION: [{"_id": vehicle_id, "registration_number": "WP CAA-1208", "vehicle_code": "QA-VAN-1"}],
    }

    async def fake_find(collection, query, fields, **kwargs):
        return rows.get(collection, [])

    monkeypatch.setattr(report_service, "_find_rows", fake_find)
    payload = await report_service.build_specialized_xlsx_report(report_type="driver_details")
    matrix = _matrix(payload, "Driver Details")
    text = "\n".join(str(value) for row in matrix for value in row if value is not None)

    assert {"Driver", "Username", "Account Status", "Device Enrollment", "Assigned Vehicle", "Completed Routes"} <= set(matrix[4])
    assert "Nimal Perera" in text
    assert "WP CAA-1208" in text
    assert "never-export" not in text
    assert "password" not in text.lower()


@pytest.mark.asyncio
async def test_route_details_report_links_consolidation_load_vehicle_and_driver(monkeypatch) -> None:
    route_id, driver_id, vehicle_id = ObjectId(), ObjectId(), ObjectId()
    rows = {
        ROUTE_PLAN_COLLECTION: [{
            "_id": route_id, "route_code": "ROUTE-001", "route_set_id": "RSET-001",
            "consolidation_id": "PCID202608291H001", "load_plan_id": "LOADPLAN-001", "load_id": "LOAD-001",
            "vehicle_id": vehicle_id, "driver_id": driver_id, "delivery_date": "2026-08-29",
            "delivery_period": "first_half", "delivery_order_ids": [ObjectId()], "status": "approved",
            "total_predicted_eta_minutes": 74.0, "average_delay_probability": 0.2, "approved_at": NOW,
            "created_at": NOW,
        }],
        ROUTE_SEGMENT_COLLECTION: [{"_id": ObjectId(), "route_plan_id": route_id, "distance_km": 8.5}],
        VEHICLE_COLLECTION: [{"_id": vehicle_id, "registration_number": "WP CAA-1208", "vehicle_code": "QA-VAN-1"}],
        USER_COLLECTION: [{"_id": driver_id, "full_name": "Nimal Perera"}],
        DELIVERY_COLLECTION: [],
    }

    async def fake_find(collection, query, fields, **kwargs):
        return rows.get(collection, [])

    monkeypatch.setattr(report_service, "_find_rows", fake_find)
    payload = await report_service.build_specialized_xlsx_report(
        report_type="route_details", consolidation_id="PCID202608291H001", route_status="approved",
    )
    matrix = _matrix(payload, "Route Details")
    headers, row = matrix[4], matrix[5]

    assert {"Route Set ID", "Consolidation ID", "Load Plan ID", "Load ID", "Vehicle", "Driver"} <= set(headers)
    assert row[headers.index("Consolidation ID")] == "PCID202608291H001"
    assert row[headers.index("Driver")] == "Nimal Perera"
