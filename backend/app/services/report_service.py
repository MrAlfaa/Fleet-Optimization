from __future__ import annotations

from datetime import date, datetime, time, timezone
from io import BytesIO
from typing import Any

from bson import ObjectId
from fastapi import HTTPException, status

from app.core.database import get_database
from app.models.alert_model import ALERT_COLLECTION
from app.models.delivery_model import DELIVERY_COLLECTION
from app.models.driver_device_model import DRIVER_DEVICE_COLLECTION
from app.models.demand_model import DEMAND_FORECAST_COLLECTION
from app.models.demand_model import (
    DEMAND_LEVELS,
    FLEET_RECOMMENDATION_COLLECTION,
    OPERATIONAL_DEMAND_DAILY_COLLECTION,
)
from app.models.inspection_model import (
    INSPECTION_COLLECTION,
    INSPECTION_PRIORITIES,
    INSPECTION_STATUSES,
)
from app.models.maintenance_model import MAINTENANCE_COLLECTION
from app.models.load_optimization_model import PARCEL_CONSOLIDATIONS_COLLECTION
from app.models.predictive_model import PREDICTION_COLLECTION
from app.models.route_model import ROUTE_PLAN_COLLECTION, ROUTE_SEGMENT_COLLECTION
from app.models.user_model import USER_COLLECTION
from app.models.vehicle_model import VEHICLE_CATEGORIES, VEHICLE_COLLECTION
from app.utils.route_utils import ROUTE_STATUSES
from app.utils.report_utils import safe_spreadsheet_value
from app.utils.sri_lanka_zones import enrich_zone_record


SECTIONS = {
    "Vehicles": (VEHICLE_COLLECTION, ["registration_number", "vehicle_code", "vehicle_type", "status", "health_status", "current_mileage", "updated_at"]),
    "Maintenance": (MAINTENANCE_COLLECTION, ["part_name", "maintenance_type", "service_date", "priority", "status", "cost"]),
    "Predictions": (PREDICTION_COLLECTION, ["vehicle_id", "part_name", "failure_probability", "risk_level", "maintenance_type_name", "predicted_rul", "rul_unit", "created_at"]),
    "Demand": (DEMAND_FORECAST_COLLECTION, ["zone_display_name", "operational_zone_id", "forecast_date", "predicted_demand", "model_type"]),
    "Routes": (ROUTE_PLAN_COLLECTION, ["route_code", "status", "total_predicted_eta_minutes", "average_delay_probability", "optimization_method", "created_at"]),
    "Deliveries": (DELIVERY_COLLECTION, ["order_code", "customer_name", "zone_id", "priority", "delivery_status", "scheduled_date"]),
    "Alerts": (ALERT_COLLECTION, ["alert_type", "severity", "title", "status", "created_at", "resolved_at"]),
}

SPECIALIZED_REPORT_TYPES = {
    "vehicle_maintenance",
    "vehicle_details",
    "daily_inspection",
    "inspection_summary",
    "demand_forecast",
    "demand_capacity",
    "route_optimization",
    "route_stop_performance",
    "consolidation_details",
    "driver_details",
    "route_details",
}
ASSESSMENT_RESULTS = {"low", "medium", "high", "critical"}
WORKFLOW_STATUSES = {"clear", "service_hold", "inspection", "maintenance"}
DELIVERY_PERIODS = {"first_half", "second_half"}
CONSOLIDATION_STATUSES = {"active", "planning", "planned", "approved", "routed", "superseded", "failed", "cancelled"}
DRIVER_ACCOUNT_STATUSES = {"active", "inactive"}
DEVICE_ENROLLMENT_STATUSES = {"pending", "enrolled", "revoked"}
MAX_REPORT_ROWS = 5000
MAX_SOURCE_ROWS = 50000

VEHICLE_FIELDS = (
    "registration_number", "vehicle_code", "vehicle_category", "fuel_type", "brand", "model",
    "vehicle_type", "year", "in_service_date", "payload_capacity_kg", "current_mileage", "status",
    "health_status", "assigned_driver_id", "retired_at", "retirement_reason", "created_at", "updated_at",
)
ASSESSMENT_FIELDS = (
    "vehicle_id", "risk_level", "failure_probability_30_days", "failure_probability_2000_km",
    "system_risks", "planning_status", "planning_window", "data_quality", "dispatch_hold",
    "recommended_action", "created_at",
)
INSPECTION_FIELDS = (
    "vehicle_id", "assessment_id", "status", "is_active", "priority", "inspection_areas",
    "assigned_to", "started_at", "completed_at", "elapsed_seconds", "checklist", "summary",
    "decision", "decision_notes", "created_at", "updated_at",
)
WORK_ORDER_FIELDS = (
    "vehicle_id", "inspection_id", "maintenance_system", "work_order_status", "status", "priority",
    "service_date", "created_at", "started_at", "completed_at", "updated_at",
)
USER_EXPORT_FIELDS = ("full_name",)
DRIVER_USER_FIELDS = (
    "full_name", "username", "email", "phone", "role", "is_active", "device_enrollment_status",
    "last_login_at", "created_at", "updated_at",
)
DRIVER_DEVICE_FIELDS = ("user_id", "status", "device_model", "enrolled_at", "last_seen_at")
DEMAND_FORECAST_FIELDS = (
    "model_run_id", "zone_id", "forecast_date", "predicted_demand", "lower_bound", "upper_bound",
)
DEMAND_RECOMMENDATION_FIELDS = (
    "model_run_id", "zone_id", "forecast_date", "predicted_demand", "demand_level",
    "estimated_required_vehicles", "recommended_action",
)
DEMAND_ACTUAL_FIELDS = ("operational_zone_id", "demand_date", "total_orders")
ROUTE_FIELDS = (
    "route_code", "vehicle_id", "driver_id", "delivery_date", "delivery_period", "delivery_order_ids",
    "optimized_stop_order", "total_predicted_eta_minutes", "average_delay_probability", "optimization_method",
    "status", "route_set_id", "consolidation_id", "load_plan_id", "load_id", "approved_at",
    "driver_assigned_at", "created_at", "dispatched_at", "completed_at",
)
ROUTE_SEGMENT_FIELDS = (
    "route_plan_id", "to_stop_index", "distance_km", "predicted_eta_minutes", "delay_probability",
    "delay_risk_level", "created_at",
)
DELIVERY_REPORT_FIELDS = (
    "order_code", "zone_id", "priority", "parcel_weight", "delivery_status", "load_planning_status",
    "active_consolidation_id", "load_profile", "created_at",
)
ROUTE_VEHICLE_FIELDS = ("registration_number", "vehicle_code")
CONSOLIDATION_FIELDS = (
    "consolidation_id", "delivery_date", "delivery_period", "status", "parcel_count", "parcel_ids",
    "parcel_points", "active_plan_id", "active_route_set_id", "cluster_method", "created_at", "updated_at",
)

VEHICLE_MAINTENANCE_HEADERS = (
    "Registration No", "Fleet Code", "Category", "Fuel Type", "Current Mileage", "Operational Status",
    "Assessment Date", "Assessment Result", "30-Day Service Risk %", "2,000-km Service Risk %",
    "Inspection Priority Areas", "Maintenance Planning Range", "Data Quality", "Service Hold",
    "Inspection Status", "Work Order Status", "Recommended Next Action",
)
VEHICLE_DETAILS_HEADERS = (
    "Registration No", "Fleet Code", "Category", "Fuel Type", "Brand", "Model", "Vehicle Type",
    "Manufacture Year", "In-Service Date", "Payload Capacity (kg)", "Current Mileage",
    "Operational Status", "Health State", "Assigned Driver", "Retired Date",
    "Retirement Reason", "Created Date", "Updated Date",
)
INSPECTION_HEADERS = (
    "Inspection Created Date", "Registration No", "Fleet Code", "Priority", "Inspection Areas", "Status",
    "Assigned Inspector", "Started Time", "Completed Time", "Elapsed Minutes", "Passed Checks",
    "Failed Checks", "Not Applicable Checks", "Summary", "Decision",
)
DEMAND_FORECAST_HEADERS = (
    "Forecast Date", "Zone", "Recorded Orders", "Forecast Orders", "Forecast Lower Bound",
    "Forecast Upper Bound", "Variance", "Demand Level", "Required Vehicle Shifts", "Recommended Action",
)
DEMAND_CAPACITY_HEADERS = (
    "Zone", "Forecast Days", "Total Forecast Orders", "Daily Average", "Peak Demand", "Peak Date",
    "High Demand Days", "Maximum Vehicle Shifts", "Capacity Recommendation",
)
ROUTE_OPTIMIZATION_HEADERS = (
    "Route Code", "Delivery Date", "Delivery Period", "Status", "Vehicle", "Driver", "Parcel Stops",
    "Total Distance (km)", "Planned ETA (minutes)", "Average Delay Risk %", "Optimization Basis",
    "Created Time", "Dispatched Time", "Completed Time",
)
ROUTE_STOP_HEADERS = (
    "Route Code", "Delivery Date", "Delivery Period", "Stop Sequence", "Order Code", "Zone", "Priority",
    "Parcel Weight (kg)", "Delivery Status", "Leg Distance (km)", "Leg ETA (minutes)",
    "Delay Risk Level", "Delay Risk %",
)
CONSOLIDATION_HEADERS = (
    "Consolidation ID", "Delivery Date", "Delivery Period", "Current Status", "Parcel Count", "Region",
    "Parcel ID", "Order Code", "Parcel Status", "Planning State", "Zone", "Parcel Weight (kg)", "Fragile",
    "Load Plan ID", "Route Set ID", "Created Time", "Updated Time",
)
DRIVER_DETAILS_HEADERS = (
    "Driver", "Username", "Email", "Phone", "Account Status", "Device Enrollment", "Device Model",
    "Last Device Seen", "Assigned Vehicle", "Active Route", "Active Route Status", "Total Routes",
    "Completed Routes", "Assigned Parcels", "Delivered Parcels", "Failed Parcels", "Last Login", "Created Date",
)
ROUTE_DETAILS_HEADERS = (
    "Route Code", "Route Set ID", "Consolidation ID", "Load Plan ID", "Load ID", "Delivery Date",
    "Delivery Period", "Status", "Vehicle", "Driver", "Parcel Stops", "Total Distance (km)",
    "Planned ETA (minutes)", "Average Delay Risk %", "Optimization Basis", "Created Time", "Approved Time",
    "Driver Assigned Time", "Dispatched Time", "Completed Time",
)

ASSESSMENT_LABELS = {
    "low": "Routine",
    "medium": "Monitor",
    "high": "Inspection Recommended",
    "critical": "Immediate Service Review",
}


def _db():
    database = get_database()
    if database is None:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="MongoDB is not configured or connected")
    return database


def _date_query(start_date: datetime | None, end_date: datetime | None) -> dict[str, Any]:
    if not start_date and not end_date:
        return {}
    value: dict[str, datetime] = {}
    if start_date:
        value["$gte"] = start_date
    if end_date:
        value["$lte"] = end_date
    return {"created_at": value}


async def collect_report_data(start_date: datetime | None = None, end_date: datetime | None = None) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {}
    for name, (collection_name, fields) in SECTIONS.items():
        projection = {field: 1 for field in fields}
        if name == "Demand":
            projection["zone_id"] = 1
        query = await _report_collection_query(name, start_date, end_date)
        rows = await _db()[collection_name].find(query, projection).sort("_id", -1).limit(5000).to_list(5000)
        cleaned = []
        for row in rows:
            if name == "Demand":
                row = enrich_zone_record(row)
            cleaned.append({field: row.get(field) for field in fields})
        result[name] = cleaned
    return result


async def _report_collection_query(
    section_name: str,
    start_date: datetime | None,
    end_date: datetime | None,
) -> dict[str, Any]:
    if section_name != "Demand":
        return _date_query(start_date, end_date)

    from app.services.demand_forecasting_service import get_active_run_id

    active_run_id = await get_active_run_id()
    if active_run_id is None:
        return {"_id": {"$exists": False}}
    query: dict[str, Any] = {"model_run_id": active_run_id}
    date_filter: dict[str, datetime] = {}
    if start_date:
        date_filter["$gte"] = start_date
    if end_date:
        date_filter["$lte"] = end_date
    if date_filter:
        query["forecast_date"] = date_filter
    return query


def _display(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).isoformat()
    return safe_spreadsheet_value("" if value is None else str(value))


def _projection(fields: tuple[str, ...]) -> dict[str, int]:
    return {field: 1 for field in fields}


async def _find_rows(
    collection_name: str,
    query: dict[str, Any],
    fields: tuple[str, ...],
    *,
    sort_field: str = "created_at",
    limit: int = MAX_SOURCE_ROWS,
) -> list[dict[str, Any]]:
    cursor = _db()[collection_name].find(query, _projection(fields)).sort(sort_field, -1).limit(limit)
    return await cursor.to_list(limit)


def _as_utc_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
    if isinstance(value, date):
        return datetime.combine(value, time.min, timezone.utc)
    if isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
            return parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed.astimezone(timezone.utc)
        except ValueError:
            return None
    return None


def _in_date_range(
    row: dict[str, Any], fields: tuple[str, ...], start_date: datetime | None, end_date: datetime | None
) -> bool:
    if start_date is None and end_date is None:
        return True
    values = [_as_utc_datetime(row.get(field)) for field in fields]
    for value in (item for item in values if item is not None):
        if start_date is not None and value < start_date:
            continue
        if end_date is not None and value > end_date:
            continue
        return True
    return False


def _event_timestamp(row: dict[str, Any]) -> float:
    for field in ("completed_at", "service_date", "updated_at", "created_at"):
        value = _as_utc_datetime(row.get(field))
        if value is not None:
            return value.timestamp()
    return 0.0


def _latest_by_vehicle(
    rows: list[dict[str, Any]], *, active_statuses: set[str] | None = None, status_field: str = "status"
) -> dict[ObjectId, dict[str, Any]]:
    grouped: dict[ObjectId, list[dict[str, Any]]] = {}
    for row in rows:
        vehicle_id = row.get("vehicle_id")
        if isinstance(vehicle_id, ObjectId):
            grouped.setdefault(vehicle_id, []).append(row)
    result: dict[ObjectId, dict[str, Any]] = {}
    for vehicle_id, candidates in grouped.items():
        def active_rank(row: dict[str, Any]) -> bool:
            if "is_active" in row:
                return row.get("is_active") is True
            return bool(active_statuses and str(row.get(status_field) or row.get("status")) in active_statuses)

        candidates.sort(
            key=lambda row: (
                active_rank(row),
                _event_timestamp(row),
            ),
            reverse=True,
        )
        result[vehicle_id] = candidates[0]
    return result


def _humanize(value: Any) -> str:
    return str(value or "").replace("_", " ").strip().title()


def _system_name(item: Any) -> str:
    if isinstance(item, dict):
        return _humanize(
            item.get("maintenance_system") or item.get("system") or item.get("name") or item.get("part_name")
        )
    return _humanize(item)


def _inspection_areas(assessment: dict[str, Any] | None, inspection: dict[str, Any] | None) -> str:
    source = (inspection or {}).get("inspection_areas") or (assessment or {}).get("system_risks") or []
    values = [value for value in (_system_name(item) for item in source) if value]
    return ", ".join(dict.fromkeys(values)) or "Routine vehicle checks"


def _planning_range(assessment: dict[str, Any] | None) -> str:
    if not assessment:
        return "Not assessed"
    window = assessment.get("planning_window") or {}
    parts: list[str] = []
    days = window.get("days") if isinstance(window, dict) else None
    km = window.get("km") if isinstance(window, dict) else None
    if isinstance(days, dict) and days.get("lower") is not None and days.get("upper") is not None:
        parts.append(f"{int(days['lower']):,}-{int(days['upper']):,} days")
    if isinstance(km, dict) and km.get("lower") is not None and km.get("upper") is not None:
        parts.append(f"{int(km['lower']):,}-{int(km['upper']):,} km")
    if parts:
        return " / ".join(parts)
    status_value = assessment.get("planning_status")
    if status_value == "no_near_term_failure_indicated":
        return "No near-term service window indicated"
    return "Not available"


def _active_work_order(record: dict[str, Any] | None) -> bool:
    if not record:
        return False
    work_status = str(record.get("work_order_status") or "")
    legacy_status = str(record.get("status") or "")
    return work_status in {"planned", "approved", "in_progress"} or (
        not work_status and legacy_status in {"scheduled", "in_progress"}
    )


def _active_inspection(record: dict[str, Any] | None) -> bool:
    if not record:
        return False
    if record.get("is_active") is True:
        return True
    return str(record.get("status") or "") in {
        "pending", "assigned", "in_progress", "completed", "decision_pending", "maintenance_required"
    }


def _workflow_status(
    vehicle: dict[str, Any],
    assessment: dict[str, Any] | None,
    inspection: dict[str, Any] | None,
    work_order: dict[str, Any] | None,
) -> str:
    if _active_work_order(work_order):
        return "maintenance"
    if _active_inspection(inspection):
        return "inspection"
    if bool((assessment or {}).get("dispatch_hold")) or vehicle.get("status") == "under_maintenance":
        return "service_hold"
    return "clear"


def _next_action(
    vehicle: dict[str, Any],
    assessment: dict[str, Any] | None,
    inspection: dict[str, Any] | None,
    work_order: dict[str, Any] | None,
) -> str:
    workflow = _workflow_status(vehicle, assessment, inspection, work_order)
    if workflow == "maintenance":
        return "Complete active maintenance work before return to service."
    if workflow == "inspection":
        return "Complete the current inspection and record the service decision."
    if workflow == "service_hold":
        return "Resolve the service hold before the next assignment."
    return {
        "critical": "Complete an engineering review before the next assignment.",
        "high": "Schedule an inspection before the next assignment.",
        "medium": "Review during the next planned service.",
        "low": "Continue the routine service schedule.",
    }.get(str((assessment or {}).get("risk_level") or ""), "Create a current vehicle assessment.")


def _probability_percent(value: Any) -> float | None:
    try:
        return round(float(value) * 100, 2) if value is not None else None
    except (TypeError, ValueError):
        return None


def _elapsed_minutes(inspection: dict[str, Any]) -> float | None:
    seconds = inspection.get("elapsed_seconds")
    if seconds is not None:
        try:
            return round(float(seconds) / 60, 1)
        except (TypeError, ValueError):
            pass
    started = _as_utc_datetime(inspection.get("started_at"))
    completed = _as_utc_datetime(inspection.get("completed_at"))
    if started is not None and completed is not None and completed >= started:
        return round((completed - started).total_seconds() / 60, 1)
    return None


def _check_counts(inspection: dict[str, Any]) -> tuple[int, int, int]:
    outcomes = [str(item.get("outcome") or "") for item in inspection.get("checklist") or [] if isinstance(item, dict)]
    return outcomes.count("passed"), outcomes.count("failed"), outcomes.count("not_applicable")


def _safe_xlsx_value(value: Any) -> Any:
    if isinstance(value, datetime):
        utc_value = value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
        return utc_value.replace(tzinfo=None)
    if isinstance(value, date):
        return value
    if isinstance(value, ObjectId):
        return safe_spreadsheet_value(str(value))
    if isinstance(value, str):
        return safe_spreadsheet_value(value)
    return value


def _filter_description(
    *,
    report_type: str,
    start_date: datetime | None,
    end_date: datetime | None,
    vehicle_id: str | None,
    vehicle_category: str | None,
    assessment_result: str | None,
    workflow_status: str | None,
    inspection_status: str | None,
    zone_id: str | None = None,
    demand_level: str | None = None,
    route_status: str | None = None,
    delivery_period: str | None = None,
    consolidation_id: str | None = None,
    consolidation_status: str | None = None,
    driver_id: str | None = None,
    driver_account_status: str | None = None,
    device_enrollment_status: str | None = None,
) -> str:
    date_scope = {
        "vehicle_maintenance": "Assessment created, inspection created/completed, or service activity date",
        "vehicle_details": "Vehicle created date",
        "daily_inspection": "Inspection created/completed date",
        "inspection_summary": "Inspection created/completed date",
        "demand_forecast": "Forecast delivery date",
        "demand_capacity": "Forecast delivery date",
        "route_optimization": "Route delivery or creation date",
        "route_stop_performance": "Route delivery or creation date",
        "consolidation_details": "Consolidation delivery or creation date",
        "driver_details": "Driver account created date",
        "route_details": "Route delivery or creation date",
    }[report_type]
    applied = [f"Date scope: {date_scope}"]
    if start_date:
        applied.append(f"start_date={start_date.date().isoformat()}")
    if end_date:
        applied.append(f"end_date={end_date.date().isoformat()}")
    for key, value in (
        ("vehicle_id", vehicle_id),
        ("vehicle_category", vehicle_category),
        ("assessment_result", assessment_result if report_type == "vehicle_maintenance" else None),
        ("workflow_status", workflow_status if report_type == "vehicle_maintenance" else None),
        ("inspection_status", inspection_status if report_type != "vehicle_details" else None),
        ("zone_id", zone_id if report_type in {"demand_forecast", "demand_capacity"} else None),
        ("demand_level", demand_level if report_type in {"demand_forecast", "demand_capacity"} else None),
        ("route_status", route_status if report_type in {"route_optimization", "route_stop_performance", "route_details"} else None),
        ("delivery_period", delivery_period if report_type in {"route_optimization", "route_stop_performance", "route_details", "consolidation_details"} else None),
        ("consolidation_id", consolidation_id if report_type in {"consolidation_details", "route_details"} else None),
        ("consolidation_status", consolidation_status if report_type == "consolidation_details" else None),
        ("driver_id", driver_id if report_type in {"driver_details", "route_details"} else None),
        ("driver_account_status", driver_account_status if report_type == "driver_details" else None),
        ("device_enrollment_status", device_enrollment_status if report_type == "driver_details" else None),
    ):
        if value:
            applied.append(f"{key}={value}")
    return "; ".join(applied)


def _validate_specialized_filters(
    *,
    report_type: str,
    start_date: datetime | None,
    end_date: datetime | None,
    vehicle_id: str | None,
    vehicle_category: str | None,
    assessment_result: str | None,
    workflow_status: str | None,
    inspection_status: str | None,
    demand_level: str | None = None,
    route_status: str | None = None,
    delivery_period: str | None = None,
    consolidation_status: str | None = None,
    driver_account_status: str | None = None,
    device_enrollment_status: str | None = None,
) -> ObjectId | None:
    if report_type not in SPECIALIZED_REPORT_TYPES:
        raise HTTPException(status_code=422, detail="Unsupported specialized report type")
    if start_date and end_date and start_date > end_date:
        raise HTTPException(status_code=422, detail="start_date must not be after end_date")
    if vehicle_category and vehicle_category not in VEHICLE_CATEGORIES:
        raise HTTPException(status_code=422, detail="Invalid vehicle_category")
    if assessment_result and assessment_result not in ASSESSMENT_RESULTS:
        raise HTTPException(status_code=422, detail="Invalid assessment_result")
    if workflow_status and workflow_status not in WORKFLOW_STATUSES:
        raise HTTPException(status_code=422, detail="Invalid workflow_status")
    if inspection_status and inspection_status not in INSPECTION_STATUSES:
        raise HTTPException(status_code=422, detail="Invalid inspection_status")
    if demand_level and demand_level not in DEMAND_LEVELS:
        raise HTTPException(status_code=422, detail="Invalid demand_level")
    if route_status and route_status not in ROUTE_STATUSES:
        raise HTTPException(status_code=422, detail="Invalid route_status")
    if delivery_period and delivery_period not in DELIVERY_PERIODS:
        raise HTTPException(status_code=422, detail="Invalid delivery_period")
    if consolidation_status and consolidation_status not in CONSOLIDATION_STATUSES:
        raise HTTPException(status_code=422, detail="Invalid consolidation_status")
    if driver_account_status and driver_account_status not in DRIVER_ACCOUNT_STATUSES:
        raise HTTPException(status_code=422, detail="Invalid driver_account_status")
    if device_enrollment_status and device_enrollment_status not in DEVICE_ENROLLMENT_STATUSES:
        raise HTTPException(status_code=422, detail="Invalid device_enrollment_status")
    if vehicle_id:
        if not ObjectId.is_valid(vehicle_id):
            raise HTTPException(status_code=422, detail="vehicle_id must be a valid ObjectId")
        return ObjectId(vehicle_id)
    return None


def _optional_object_id(value: str | None, field_name: str) -> ObjectId | None:
    if not value:
        return None
    if not ObjectId.is_valid(value):
        raise HTTPException(status_code=422, detail=f"{field_name} must be a valid ObjectId")
    return ObjectId(value)


async def _report_sources(
    *,
    vehicle_oid: ObjectId | None,
    vehicle_category: str | None,
    start_date: datetime | None,
    end_date: datetime | None,
    include_assessments: bool,
    include_inspections: bool,
    include_work_orders: bool,
) -> tuple[
    list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], dict[str, str]
]:
    vehicle_query: dict[str, Any] = {}
    if vehicle_oid is not None:
        vehicle_query["_id"] = vehicle_oid
    if vehicle_category:
        vehicle_query["vehicle_category"] = vehicle_category
    vehicles = await _find_rows(VEHICLE_COLLECTION, vehicle_query, VEHICLE_FIELDS, limit=MAX_REPORT_ROWS)
    vehicle_ids = [row["_id"] for row in vehicles]
    if not vehicle_ids:
        return [], [], [], [], {}
    related_query = {"vehicle_id": {"$in": vehicle_ids}}
    assessments = (
        await _find_rows(PREDICTION_COLLECTION, related_query, ASSESSMENT_FIELDS) if include_assessments else []
    )
    inspections = (
        await _find_rows(INSPECTION_COLLECTION, related_query, INSPECTION_FIELDS) if include_inspections else []
    )
    work_orders = (
        await _find_rows(MAINTENANCE_COLLECTION, related_query, WORK_ORDER_FIELDS, sort_field="service_date")
        if include_work_orders else []
    )
    user_ids = {
        value for value in [*(row.get("assigned_driver_id") for row in vehicles), *(row.get("assigned_to") for row in inspections)]
        if isinstance(value, ObjectId)
    }
    users = (
        await _find_rows(USER_COLLECTION, {"_id": {"$in": list(user_ids)}}, USER_EXPORT_FIELDS, limit=len(user_ids))
        if user_ids else []
    )
    user_names = {str(row["_id"]): str(row.get("full_name") or "") for row in users}
    return vehicles, assessments, inspections, work_orders, user_names


def _vehicle_maintenance_rows(
    vehicles: list[dict[str, Any]],
    assessments: list[dict[str, Any]],
    inspections: list[dict[str, Any]],
    work_orders: list[dict[str, Any]],
    *,
    start_date: datetime | None,
    end_date: datetime | None,
    assessment_result: str | None,
    workflow_status: str | None,
    inspection_status: str | None,
) -> list[list[Any]]:
    assessments = [row for row in assessments if _in_date_range(row, ("created_at",), start_date, end_date)]
    inspections = [
        row for row in inspections
        if _in_date_range(row, ("created_at", "completed_at"), start_date, end_date)
    ]
    work_orders = [
        row for row in work_orders
        if _in_date_range(row, ("service_date", "created_at", "completed_at"), start_date, end_date)
    ]
    latest_assessments = _latest_by_vehicle(assessments)
    latest_inspections = _latest_by_vehicle(
        inspections,
        active_statuses={"pending", "assigned", "in_progress", "completed", "decision_pending", "maintenance_required"},
    )
    latest_work = _latest_by_vehicle(
        work_orders, active_statuses={"planned", "approved", "in_progress"}, status_field="work_order_status"
    )
    rows: list[list[Any]] = []
    date_filtered = start_date is not None or end_date is not None
    for vehicle in vehicles:
        vehicle_id = vehicle["_id"]
        assessment = latest_assessments.get(vehicle_id)
        inspection = latest_inspections.get(vehicle_id)
        work_order = latest_work.get(vehicle_id)
        if date_filtered and not any((assessment, inspection, work_order)):
            continue
        if assessment_result and str((assessment or {}).get("risk_level") or "") != assessment_result:
            continue
        if inspection_status and str((inspection or {}).get("status") or "") != inspection_status:
            continue
        current_workflow = _workflow_status(vehicle, assessment, inspection, work_order)
        if workflow_status and current_workflow != workflow_status:
            continue
        risk_level = str((assessment or {}).get("risk_level") or "")
        rows.append([
            vehicle.get("registration_number"), vehicle.get("vehicle_code"), _humanize(vehicle.get("vehicle_category")),
            _humanize(vehicle.get("fuel_type")), vehicle.get("current_mileage"), _humanize(vehicle.get("status")),
            (assessment or {}).get("created_at"), ASSESSMENT_LABELS.get(risk_level, "Not Assessed"),
            _probability_percent((assessment or {}).get("failure_probability_30_days")),
            _probability_percent((assessment or {}).get("failure_probability_2000_km")),
            _inspection_areas(assessment, inspection), _planning_range(assessment),
            _humanize((assessment or {}).get("data_quality")) or "Not Available",
            "Yes" if bool((assessment or {}).get("dispatch_hold")) or current_workflow != "clear" else "No",
            _humanize((inspection or {}).get("status")) or "No Inspection",
            _humanize((work_order or {}).get("work_order_status") or (work_order or {}).get("status")) or "No Work Order",
            _next_action(vehicle, assessment, inspection, work_order),
        ])
    return rows[:MAX_REPORT_ROWS]


def _vehicle_detail_rows(
    vehicles: list[dict[str, Any]], user_names: dict[str, str], start_date: datetime | None, end_date: datetime | None
) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for vehicle in vehicles:
        if not _in_date_range(vehicle, ("created_at",), start_date, end_date):
            continue
        driver_id = vehicle.get("assigned_driver_id")
        rows.append([
            vehicle.get("registration_number"), vehicle.get("vehicle_code"), _humanize(vehicle.get("vehicle_category")),
            _humanize(vehicle.get("fuel_type")), vehicle.get("brand"), vehicle.get("model"),
            _humanize(vehicle.get("vehicle_type")), vehicle.get("year"), vehicle.get("in_service_date"),
            vehicle.get("payload_capacity_kg"), vehicle.get("current_mileage"), _humanize(vehicle.get("status")),
            _humanize(vehicle.get("health_status")), user_names.get(str(driver_id), "") if driver_id else "",
            vehicle.get("retired_at"),
            vehicle.get("retirement_reason"), vehicle.get("created_at"), vehicle.get("updated_at"),
        ])
    return rows[:MAX_REPORT_ROWS]


def _inspection_detail_rows(
    vehicles: list[dict[str, Any]],
    inspections: list[dict[str, Any]],
    user_names: dict[str, str],
    *,
    start_date: datetime | None,
    end_date: datetime | None,
    inspection_status: str | None,
) -> tuple[list[list[Any]], list[dict[str, Any]]]:
    vehicle_map = {row["_id"]: row for row in vehicles}
    filtered = [
        row for row in inspections
        if row.get("vehicle_id") in vehicle_map
        and _in_date_range(row, ("created_at", "completed_at"), start_date, end_date)
        and (not inspection_status or row.get("status") == inspection_status)
    ][:MAX_REPORT_ROWS]
    rows: list[list[Any]] = []
    for inspection in filtered:
        vehicle = vehicle_map[inspection["vehicle_id"]]
        passed, failed, not_applicable = _check_counts(inspection)
        assignee_id = inspection.get("assigned_to")
        rows.append([
            inspection.get("created_at"), vehicle.get("registration_number"), vehicle.get("vehicle_code"),
            _humanize(inspection.get("priority")), _inspection_areas(None, inspection),
            _humanize(inspection.get("status")), user_names.get(str(assignee_id), "") if assignee_id else "",
            inspection.get("started_at"),
            inspection.get("completed_at"), _elapsed_minutes(inspection), passed, failed, not_applicable,
            inspection.get("summary") or inspection.get("decision_notes") or "",
            _humanize(inspection.get("decision")) or "Pending Decision",
        ])
    return rows, filtered


async def _active_demand_run_id() -> ObjectId | None:
    from app.services.demand_forecasting_service import get_active_run_id

    return await get_active_run_id()


def _date_key(value: Any) -> str:
    parsed = _as_utc_datetime(value)
    return parsed.date().isoformat() if parsed else str(value or "")[:10]


def _zone_name(zone_id: Any) -> str:
    if not zone_id:
        return "Unassigned Zone"
    enriched = enrich_zone_record({"zone_id": str(zone_id)})
    return str(enriched.get("zone_display_name") or enriched.get("zone_name") or _humanize(zone_id))


def _route_basis(value: Any) -> str:
    method = str(value or "").lower()
    return "Traffic-aware optimized route" if "predictive" in method else "Distance-optimized route"


async def _demand_report_rows(
    *,
    summary: bool,
    start_date: datetime | None,
    end_date: datetime | None,
    zone_id: str | None,
    demand_level: str | None,
) -> list[list[Any]]:
    run_id = await _active_demand_run_id()
    if run_id is None:
        return []
    forecasts = await _find_rows(
        DEMAND_FORECAST_COLLECTION, {"model_run_id": run_id}, DEMAND_FORECAST_FIELDS,
        sort_field="forecast_date", limit=MAX_SOURCE_ROWS,
    )
    recommendations = await _find_rows(
        FLEET_RECOMMENDATION_COLLECTION, {"model_run_id": run_id}, DEMAND_RECOMMENDATION_FIELDS,
        sort_field="forecast_date", limit=MAX_SOURCE_ROWS,
    )
    actuals = await _find_rows(
        OPERATIONAL_DEMAND_DAILY_COLLECTION, {}, DEMAND_ACTUAL_FIELDS,
        sort_field="demand_date", limit=MAX_SOURCE_ROWS,
    )
    recommendation_map = {
        (str(row.get("zone_id")), _date_key(row.get("forecast_date"))): row for row in recommendations
    }
    actual_map = {
        (str(row.get("operational_zone_id")), _date_key(row.get("demand_date"))): row.get("total_orders")
        for row in actuals
    }
    detail: list[dict[str, Any]] = []
    for forecast in forecasts:
        current_zone = str(forecast.get("zone_id") or "")
        current_date = _as_utc_datetime(forecast.get("forecast_date"))
        recommendation = recommendation_map.get((current_zone, _date_key(forecast.get("forecast_date"))), {})
        if zone_id and current_zone != zone_id:
            continue
        if demand_level and str(recommendation.get("demand_level") or "") != demand_level:
            continue
        if not _in_date_range({"forecast_date": current_date}, ("forecast_date",), start_date, end_date):
            continue
        predicted = float(forecast.get("predicted_demand") or 0)
        actual = actual_map.get((current_zone, _date_key(current_date)))
        detail.append({
            "zone_id": current_zone,
            "date": current_date,
            "predicted": predicted,
            "actual": actual,
            "lower": forecast.get("lower_bound"),
            "upper": forecast.get("upper_bound"),
            "variance": round(float(actual) - predicted, 2) if actual is not None else None,
            "level": str(recommendation.get("demand_level") or ""),
            "vehicles": recommendation.get("estimated_required_vehicles"),
            "action": recommendation.get("recommended_action") or "Review planned delivery capacity",
        })
    if not summary:
        return [[
            row["date"], _zone_name(row["zone_id"]), row["actual"], row["predicted"], row["lower"], row["upper"],
            row["variance"], _humanize(row["level"]), row["vehicles"], row["action"],
        ] for row in detail[:MAX_REPORT_ROWS]]

    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in detail:
        grouped.setdefault(row["zone_id"], []).append(row)
    rows: list[list[Any]] = []
    for current_zone, values in sorted(grouped.items(), key=lambda item: _zone_name(item[0])):
        peak = max(values, key=lambda item: item["predicted"])
        actions = [str(item["action"]) for item in values if item.get("action")]
        rows.append([
            _zone_name(current_zone), len(values), round(sum(item["predicted"] for item in values), 2),
            round(sum(item["predicted"] for item in values) / len(values), 2), peak["predicted"], peak["date"],
            sum(item["level"] in {"high", "very_high"} for item in values),
            max((int(item["vehicles"] or 0) for item in values), default=0),
            actions[0] if actions else "Use planned delivery capacity",
        ])
    return rows[:MAX_REPORT_ROWS]


async def _route_report_rows(
    *,
    stop_detail: bool,
    start_date: datetime | None,
    end_date: datetime | None,
    vehicle_oid: ObjectId | None,
    route_status: str | None,
    delivery_period: str | None,
) -> list[list[Any]]:
    route_query: dict[str, Any] = {"vehicle_id": vehicle_oid} if vehicle_oid else {}
    routes = await _find_rows(ROUTE_PLAN_COLLECTION, route_query, ROUTE_FIELDS, limit=MAX_SOURCE_ROWS)
    routes = [
        route for route in routes
        if (not route_status or route.get("status") == route_status)
        and (not delivery_period or route.get("delivery_period") == delivery_period)
        and _in_date_range(route, ("delivery_date", "created_at"), start_date, end_date)
    ][:MAX_REPORT_ROWS]
    if not routes:
        return []
    route_ids = [route["_id"] for route in routes]
    vehicle_ids = [route.get("vehicle_id") for route in routes if isinstance(route.get("vehicle_id"), ObjectId)]
    driver_ids = [route.get("driver_id") for route in routes if isinstance(route.get("driver_id"), ObjectId)]
    order_ids = [order_id for route in routes for order_id in route.get("delivery_order_ids") or [] if isinstance(order_id, ObjectId)]
    segments = await _find_rows(
        ROUTE_SEGMENT_COLLECTION, {"route_plan_id": {"$in": route_ids}}, ROUTE_SEGMENT_FIELDS, limit=MAX_SOURCE_ROWS,
    )
    vehicles = await _find_rows(
        VEHICLE_COLLECTION, {"_id": {"$in": vehicle_ids}}, ROUTE_VEHICLE_FIELDS, limit=len(vehicle_ids) or 1,
    ) if vehicle_ids else []
    users = await _find_rows(
        USER_COLLECTION, {"_id": {"$in": driver_ids}}, USER_EXPORT_FIELDS, limit=len(driver_ids) or 1,
    ) if driver_ids else []
    deliveries = await _find_rows(
        DELIVERY_COLLECTION, {"_id": {"$in": order_ids}}, DELIVERY_REPORT_FIELDS, limit=len(order_ids) or 1,
    ) if order_ids else []
    vehicle_map = {
        row["_id"]: " · ".join(filter(None, [str(row.get("registration_number") or ""), str(row.get("vehicle_code") or "")]))
        for row in vehicles
    }
    user_map = {row["_id"]: str(row.get("full_name") or "") for row in users}
    delivery_map = {str(row["_id"]): row for row in deliveries}
    segments_by_route: dict[ObjectId, list[dict[str, Any]]] = {}
    for segment in segments:
        route_id = segment.get("route_plan_id")
        if isinstance(route_id, ObjectId):
            segments_by_route.setdefault(route_id, []).append(segment)

    rows: list[list[Any]] = []
    for route in routes:
        route_segments = segments_by_route.get(route["_id"], [])
        if not stop_detail:
            rows.append([
                route.get("route_code"), route.get("delivery_date"), _humanize(route.get("delivery_period")),
                _humanize(route.get("status")), vehicle_map.get(route.get("vehicle_id"), "Unassigned"),
                user_map.get(route.get("driver_id"), "Unassigned"), len(route.get("delivery_order_ids") or []),
                round(sum(float(item.get("distance_km") or 0) for item in route_segments), 2),
                route.get("total_predicted_eta_minutes"), _probability_percent(route.get("average_delay_probability")),
                _route_basis(route.get("optimization_method")), route.get("created_at"), route.get("dispatched_at"),
                route.get("completed_at"),
            ])
            continue
        segment_by_stop = {int(item.get("to_stop_index") or -1): item for item in route_segments}
        delivery_stops = [item for item in route.get("optimized_stop_order") or [] if item.get("type") == "delivery"]
        for stop in delivery_stops:
            delivery = delivery_map.get(str(stop.get("order_id")), {})
            segment = segment_by_stop.get(int(stop.get("point_index") or -1), {})
            rows.append([
                route.get("route_code"), route.get("delivery_date"), _humanize(route.get("delivery_period")),
                stop.get("sequence"), delivery.get("order_code") or stop.get("order_code"),
                _zone_name(delivery.get("zone_id") or stop.get("zone_id")), _humanize(delivery.get("priority") or stop.get("priority")),
                delivery.get("parcel_weight") or stop.get("parcel_weight"), _humanize(delivery.get("delivery_status")),
                segment.get("distance_km"), segment.get("predicted_eta_minutes"),
                _humanize(segment.get("delay_risk_level")), _probability_percent(segment.get("delay_probability")),
            ])
    return rows[:MAX_REPORT_ROWS]


async def _consolidation_report_rows(
    *,
    start_date: datetime | None,
    end_date: datetime | None,
    consolidation_id: str | None,
    consolidation_status: str | None,
    delivery_period: str | None,
) -> list[list[Any]]:
    query: dict[str, Any] = {}
    if consolidation_id:
        query["consolidation_id"] = consolidation_id
    if consolidation_status:
        query["status"] = consolidation_status
    if delivery_period:
        query["delivery_period"] = delivery_period
    consolidations = await _find_rows(
        PARCEL_CONSOLIDATIONS_COLLECTION, query, CONSOLIDATION_FIELDS, limit=MAX_REPORT_ROWS,
    )
    consolidations = [
        row for row in consolidations
        if (not consolidation_id or str(row.get("consolidation_id") or "") == consolidation_id)
        and (not consolidation_status or row.get("status") == consolidation_status)
        and (not delivery_period or row.get("delivery_period") == delivery_period)
        and _in_date_range(row, ("delivery_date", "created_at"), start_date, end_date)
    ]
    parcel_ids = {
        str(parcel_id)
        for consolidation in consolidations
        for parcel_id in consolidation.get("parcel_ids") or []
        if str(parcel_id)
    }
    object_ids = [ObjectId(value) for value in parcel_ids if ObjectId.is_valid(value)]
    deliveries = await _find_rows(
        DELIVERY_COLLECTION, {"_id": {"$in": object_ids}}, DELIVERY_REPORT_FIELDS,
        limit=max(len(object_ids), 1),
    ) if object_ids else []
    delivery_map = {str(row.get("_id")): row for row in deliveries}

    rows: list[list[Any]] = []
    for consolidation in consolidations:
        point_map = {
            str(point.get("parcel_id")): point
            for point in consolidation.get("parcel_points") or []
        }
        for parcel_id in consolidation.get("parcel_ids") or []:
            parcel_key = str(parcel_id)
            delivery = delivery_map.get(parcel_key, {})
            point = point_map.get(parcel_key, {})
            load_profile = delivery.get("load_profile") or {}
            rows.append([
                consolidation.get("consolidation_id"), consolidation.get("delivery_date"),
                _humanize(consolidation.get("delivery_period")), _humanize(consolidation.get("status")),
                consolidation.get("parcel_count"), point.get("region_id"), parcel_key,
                delivery.get("order_code") or point.get("order_code") or parcel_key,
                _humanize(delivery.get("delivery_status")), _humanize(delivery.get("load_planning_status")),
                _zone_name(delivery.get("zone_id") or point.get("zone_id")), delivery.get("parcel_weight"),
                "Yes" if load_profile.get("fragile") is True else "No",
                consolidation.get("active_plan_id"), consolidation.get("active_route_set_id"),
                consolidation.get("created_at"), consolidation.get("updated_at"),
            ])
    return rows[:MAX_REPORT_ROWS]


async def _driver_detail_rows(
    *,
    start_date: datetime | None,
    end_date: datetime | None,
    driver_oid: ObjectId | None,
    driver_account_status: str | None,
    device_enrollment_status: str | None,
) -> list[list[Any]]:
    user_query: dict[str, Any] = {"role": "driver"}
    if driver_oid:
        user_query["_id"] = driver_oid
    drivers = await _find_rows(USER_COLLECTION, user_query, DRIVER_USER_FIELDS, limit=MAX_REPORT_ROWS)
    drivers = [
        driver for driver in drivers
        if driver.get("role") == "driver"
        and (not driver_oid or driver.get("_id") == driver_oid)
        and (not driver_account_status or ("active" if driver.get("is_active", True) else "inactive") == driver_account_status)
        and (not device_enrollment_status or driver.get("device_enrollment_status") == device_enrollment_status)
        and _in_date_range(driver, ("created_at",), start_date, end_date)
    ]
    driver_ids = [driver["_id"] for driver in drivers if isinstance(driver.get("_id"), ObjectId)]
    if not driver_ids:
        return []
    devices = await _find_rows(
        DRIVER_DEVICE_COLLECTION, {"user_id": {"$in": driver_ids}}, DRIVER_DEVICE_FIELDS,
        sort_field="last_seen_at", limit=MAX_SOURCE_ROWS,
    )
    routes = await _find_rows(
        ROUTE_PLAN_COLLECTION, {"driver_id": {"$in": driver_ids}}, ROUTE_FIELDS, limit=MAX_SOURCE_ROWS,
    )
    vehicle_ids = [route.get("vehicle_id") for route in routes if isinstance(route.get("vehicle_id"), ObjectId)]
    vehicles = await _find_rows(
        VEHICLE_COLLECTION, {"_id": {"$in": vehicle_ids}}, ROUTE_VEHICLE_FIELDS,
        limit=max(len(vehicle_ids), 1),
    ) if vehicle_ids else []
    order_ids = [
        order_id for route in routes for order_id in route.get("delivery_order_ids") or []
        if isinstance(order_id, ObjectId)
    ]
    deliveries = await _find_rows(
        DELIVERY_COLLECTION, {"_id": {"$in": order_ids}}, DELIVERY_REPORT_FIELDS,
        limit=max(len(order_ids), 1),
    ) if order_ids else []
    delivery_map = {str(row.get("_id")): row for row in deliveries}
    vehicle_map = {
        row["_id"]: " · ".join(filter(None, [str(row.get("registration_number") or ""), str(row.get("vehicle_code") or "")]))
        for row in vehicles
    }
    devices_by_driver: dict[ObjectId, list[dict[str, Any]]] = {}
    for device in devices:
        user_id = device.get("user_id")
        if isinstance(user_id, ObjectId):
            devices_by_driver.setdefault(user_id, []).append(device)
    routes_by_driver: dict[ObjectId, list[dict[str, Any]]] = {}
    for route in routes:
        driver_id = route.get("driver_id")
        if isinstance(driver_id, ObjectId):
            routes_by_driver.setdefault(driver_id, []).append(route)

    result: list[list[Any]] = []
    for driver in sorted(drivers, key=lambda row: str(row.get("full_name") or "")):
        driver_id = driver["_id"]
        driver_devices = sorted(
            devices_by_driver.get(driver_id, []), key=lambda row: _event_timestamp(row), reverse=True,
        )
        active_device = next((item for item in driver_devices if item.get("status") == "active"), None)
        driver_routes = routes_by_driver.get(driver_id, [])
        ordered_routes = sorted(driver_routes, key=_event_timestamp, reverse=True)
        active_route = next(
            (route for route in ordered_routes if route.get("status") in {"dispatched", "in_progress"}),
            None,
        )
        latest_route = ordered_routes[0] if ordered_routes else None
        parcel_ids = [str(order_id) for route in driver_routes for order_id in route.get("delivery_order_ids") or []]
        statuses = [str(delivery_map.get(parcel_id, {}).get("delivery_status") or "") for parcel_id in parcel_ids]
        result.append([
            driver.get("full_name"), driver.get("username"), driver.get("email"), driver.get("phone"),
            "Active" if driver.get("is_active", True) else "Inactive",
            _humanize(driver.get("device_enrollment_status")),
            (active_device or (driver_devices[0] if driver_devices else {})).get("device_model"),
            (active_device or (driver_devices[0] if driver_devices else {})).get("last_seen_at"),
            vehicle_map.get((active_route or latest_route or {}).get("vehicle_id"), "Unassigned"),
            active_route.get("route_code") if active_route else "None",
            _humanize(active_route.get("status")) if active_route else "None",
            len(driver_routes), sum(route.get("status") == "completed" for route in driver_routes), len(parcel_ids),
            statuses.count("delivered"), statuses.count("failed"), driver.get("last_login_at"), driver.get("created_at"),
        ])
    return result[:MAX_REPORT_ROWS]


async def _route_details_report_rows(
    *,
    start_date: datetime | None,
    end_date: datetime | None,
    vehicle_oid: ObjectId | None,
    driver_oid: ObjectId | None,
    route_status: str | None,
    delivery_period: str | None,
    consolidation_id: str | None,
) -> list[list[Any]]:
    query: dict[str, Any] = {}
    if vehicle_oid:
        query["vehicle_id"] = vehicle_oid
    if driver_oid:
        query["driver_id"] = driver_oid
    if consolidation_id:
        query["consolidation_id"] = consolidation_id
    routes = await _find_rows(ROUTE_PLAN_COLLECTION, query, ROUTE_FIELDS, limit=MAX_SOURCE_ROWS)
    routes = [
        route for route in routes
        if (not vehicle_oid or route.get("vehicle_id") == vehicle_oid)
        and (not driver_oid or route.get("driver_id") == driver_oid)
        and (not route_status or route.get("status") == route_status)
        and (not delivery_period or route.get("delivery_period") == delivery_period)
        and (not consolidation_id or route.get("consolidation_id") == consolidation_id)
        and _in_date_range(route, ("delivery_date", "created_at"), start_date, end_date)
    ][:MAX_REPORT_ROWS]
    if not routes:
        return []
    route_ids = [route["_id"] for route in routes]
    vehicle_ids = [route.get("vehicle_id") for route in routes if isinstance(route.get("vehicle_id"), ObjectId)]
    driver_ids = [route.get("driver_id") for route in routes if isinstance(route.get("driver_id"), ObjectId)]
    segments = await _find_rows(
        ROUTE_SEGMENT_COLLECTION, {"route_plan_id": {"$in": route_ids}}, ROUTE_SEGMENT_FIELDS, limit=MAX_SOURCE_ROWS,
    )
    vehicles = await _find_rows(
        VEHICLE_COLLECTION, {"_id": {"$in": vehicle_ids}}, ROUTE_VEHICLE_FIELDS, limit=max(len(vehicle_ids), 1),
    ) if vehicle_ids else []
    users = await _find_rows(
        USER_COLLECTION, {"_id": {"$in": driver_ids}}, USER_EXPORT_FIELDS, limit=max(len(driver_ids), 1),
    ) if driver_ids else []
    vehicle_map = {
        row["_id"]: " · ".join(filter(None, [str(row.get("registration_number") or ""), str(row.get("vehicle_code") or "")]))
        for row in vehicles
    }
    user_map = {row["_id"]: str(row.get("full_name") or "") for row in users}
    distance_by_route: dict[ObjectId, float] = {}
    for segment in segments:
        route_id = segment.get("route_plan_id")
        if isinstance(route_id, ObjectId):
            distance_by_route[route_id] = distance_by_route.get(route_id, 0.0) + float(segment.get("distance_km") or 0)
    return [[
        route.get("route_code"), route.get("route_set_id"), route.get("consolidation_id"),
        route.get("load_plan_id"), route.get("load_id"), route.get("delivery_date"),
        _humanize(route.get("delivery_period")), _humanize(route.get("status")),
        vehicle_map.get(route.get("vehicle_id"), "Unassigned"), user_map.get(route.get("driver_id"), "Unassigned"),
        len(route.get("delivery_order_ids") or []), round(distance_by_route.get(route["_id"], 0.0), 2),
        route.get("total_predicted_eta_minutes"), _probability_percent(route.get("average_delay_probability")),
        _route_basis(route.get("optimization_method")), route.get("created_at"), route.get("approved_at"),
        route.get("driver_assigned_at"), route.get("dispatched_at"), route.get("completed_at"),
    ] for route in routes]


def _style_sheet(sheet: Any, header_row: int, *, auto_filter: bool = True) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.fill = PatternFill("solid", fgColor="173B8F")
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2763C4")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = f"A{header_row + 1}"
    if auto_filter and sheet.max_row >= header_row:
        sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in column]
        sheet.column_dimensions[column[0].column_letter].width = min(max(max(map(len, values), default=0) + 2, 12), 42)
    headers = {cell.column: str(cell.value or "") for cell in sheet[header_row]}
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
            elif headers.get(cell.column, "").endswith("Risk %") and cell.value is not None:
                cell.number_format = "0.00"
            elif headers.get(cell.column) in {
                "Current Mileage", "Payload Capacity (kg)", "Elapsed Minutes", "Total Distance (km)",
                "Planned ETA (minutes)", "Parcel Weight (kg)", "Leg Distance (km)", "Leg ETA (minutes)",
            } and cell.value is not None:
                cell.number_format = "#,##0.00"


def _append_safe(sheet: Any, values: tuple[Any, ...] | list[Any]) -> None:
    sheet.append([_safe_xlsx_value(value) for value in values])


def _create_detail_sheet(
    workbook: Any,
    *,
    name: str,
    title: str,
    filters: str,
    headers: tuple[str, ...],
    rows: list[list[Any]],
) -> Any:
    sheet = workbook.create_sheet()
    sheet.title = name
    _append_safe(sheet, (title,))
    _append_safe(sheet, ("Generated", datetime.now(timezone.utc)))
    _append_safe(sheet, ("Applied Filters", filters))
    sheet.append([])
    _append_safe(sheet, headers)
    for row in rows:
        _append_safe(sheet, row)
    _style_sheet(sheet, 5)
    return sheet


async def build_specialized_xlsx_report(
    *,
    report_type: str,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
    vehicle_id: str | None = None,
    vehicle_category: str | None = None,
    assessment_result: str | None = None,
    workflow_status: str | None = None,
    inspection_status: str | None = None,
    zone_id: str | None = None,
    demand_level: str | None = None,
    route_status: str | None = None,
    delivery_period: str | None = None,
    consolidation_id: str | None = None,
    consolidation_status: str | None = None,
    driver_id: str | None = None,
    driver_account_status: str | None = None,
    device_enrollment_status: str | None = None,
) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise HTTPException(status_code=503, detail="XLSX report dependency is unavailable") from exc

    vehicle_oid = _validate_specialized_filters(
        report_type=report_type, start_date=start_date, end_date=end_date, vehicle_id=vehicle_id,
        vehicle_category=vehicle_category, assessment_result=assessment_result,
        workflow_status=workflow_status, inspection_status=inspection_status, demand_level=demand_level,
        route_status=route_status, delivery_period=delivery_period, consolidation_status=consolidation_status,
        driver_account_status=driver_account_status, device_enrollment_status=device_enrollment_status,
    )
    driver_oid = _optional_object_id(driver_id, "driver_id")
    vehicle_report = report_type in {"vehicle_maintenance", "vehicle_details", "daily_inspection", "inspection_summary"}
    if vehicle_report:
        include_assessments = report_type == "vehicle_maintenance"
        include_inspections = report_type in {"vehicle_maintenance", "daily_inspection", "inspection_summary"}
        include_work_orders = report_type == "vehicle_maintenance"
        vehicles, assessments, inspections, work_orders, user_names = await _report_sources(
            vehicle_oid=vehicle_oid, vehicle_category=vehicle_category, start_date=start_date, end_date=end_date,
            include_assessments=include_assessments, include_inspections=include_inspections,
            include_work_orders=include_work_orders,
        )
    else:
        vehicles, assessments, inspections, work_orders, user_names = [], [], [], [], {}
    filters = _filter_description(
        report_type=report_type, start_date=start_date, end_date=end_date, vehicle_id=vehicle_id,
        vehicle_category=vehicle_category, assessment_result=assessment_result,
        workflow_status=workflow_status, inspection_status=inspection_status, zone_id=zone_id,
        demand_level=demand_level, route_status=route_status, delivery_period=delivery_period,
        consolidation_id=consolidation_id, consolidation_status=consolidation_status, driver_id=driver_id,
        driver_account_status=driver_account_status, device_enrollment_status=device_enrollment_status,
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    if report_type == "vehicle_maintenance":
        rows = _vehicle_maintenance_rows(
            vehicles, assessments, inspections, work_orders, start_date=start_date, end_date=end_date,
            assessment_result=assessment_result, workflow_status=workflow_status,
            inspection_status=inspection_status,
        )
        _create_detail_sheet(
            workbook, name="Vehicle Maintenance", title="Vehicle Maintenance Report", filters=filters,
            headers=VEHICLE_MAINTENANCE_HEADERS, rows=rows,
        )
    elif report_type == "vehicle_details":
        rows = _vehicle_detail_rows(vehicles, user_names, start_date, end_date)
        _create_detail_sheet(
            workbook, name="Vehicle Details", title="Vehicle Details Report", filters=filters,
            headers=VEHICLE_DETAILS_HEADERS, rows=rows,
        )
    elif report_type in {"daily_inspection", "inspection_summary"}:
        rows, filtered_inspections = _inspection_detail_rows(
            vehicles, inspections, user_names, start_date=start_date, end_date=end_date,
            inspection_status=inspection_status,
        )
        if report_type == "daily_inspection":
            _create_detail_sheet(
                workbook, name="Daily Inspection", title="Daily Inspection Report", filters=filters,
                headers=INSPECTION_HEADERS, rows=rows,
            )
        else:
            summary = workbook.create_sheet("Summary")
            _append_safe(summary, ("Inspection Summary Report",))
            _append_safe(summary, ("Generated", datetime.now(timezone.utc)))
            _append_safe(summary, ("Applied Filters", filters))
            summary.append([])
            _append_safe(summary, ("Status", "Count"))
            status_counts = {status_value: 0 for status_value in INSPECTION_STATUSES}
            priority_counts = {priority: 0 for priority in INSPECTION_PRIORITIES}
            for inspection in filtered_inspections:
                status_value = str(inspection.get("status") or "")
                priority = str(inspection.get("priority") or "")
                if status_value in status_counts:
                    status_counts[status_value] += 1
                if priority in priority_counts:
                    priority_counts[priority] += 1
            status_order = (
                "pending", "assigned", "in_progress", "decision_pending", "maintenance_required",
                "returned_to_service", "cancelled", "completed",
            )
            for status_value in status_order:
                _append_safe(summary, (_humanize(status_value), status_counts[status_value]))
            summary.append([])
            _append_safe(summary, ("Priority", "Count"))
            for priority in ("critical", "high", "medium", "low"):
                _append_safe(summary, (_humanize(priority), priority_counts[priority]))
            _style_sheet(summary, 5, auto_filter=False)
            detail = workbook.create_sheet("Inspection Details")
            _append_safe(detail, ("Inspection Details",))
            _append_safe(detail, ("Generated", datetime.now(timezone.utc)))
            _append_safe(detail, ("Applied Filters", filters))
            detail.append([])
            _append_safe(detail, INSPECTION_HEADERS)
            for row in rows:
                _append_safe(detail, row)
            _style_sheet(detail, 5)
    elif report_type in {"demand_forecast", "demand_capacity"}:
        rows = await _demand_report_rows(
            summary=report_type == "demand_capacity", start_date=start_date, end_date=end_date,
            zone_id=zone_id, demand_level=demand_level,
        )
        if report_type == "demand_forecast":
            _create_detail_sheet(
                workbook, name="Demand Forecast", title="Zone Demand Forecast Report", filters=filters,
                headers=DEMAND_FORECAST_HEADERS, rows=rows,
            )
        else:
            _create_detail_sheet(
                workbook, name="Capacity Summary", title="Demand Capacity Summary Report", filters=filters,
                headers=DEMAND_CAPACITY_HEADERS, rows=rows,
            )
    elif report_type in {"route_optimization", "route_stop_performance"}:
        rows = await _route_report_rows(
            stop_detail=report_type == "route_stop_performance", start_date=start_date, end_date=end_date,
            vehicle_oid=vehicle_oid, route_status=route_status, delivery_period=delivery_period,
        )
        if report_type == "route_optimization":
            _create_detail_sheet(
                workbook, name="Route Optimization", title="Route Optimization Report", filters=filters,
                headers=ROUTE_OPTIMIZATION_HEADERS, rows=rows,
            )
        else:
            _create_detail_sheet(
                workbook, name="Route Stops", title="Route Stop Status Report", filters=filters,
                headers=ROUTE_STOP_HEADERS, rows=rows,
            )
    elif report_type == "consolidation_details":
        rows = await _consolidation_report_rows(
            start_date=start_date, end_date=end_date, consolidation_id=consolidation_id,
            consolidation_status=consolidation_status, delivery_period=delivery_period,
        )
        _create_detail_sheet(
            workbook, name="Consolidations", title="Parcel Consolidation Report", filters=filters,
            headers=CONSOLIDATION_HEADERS, rows=rows,
        )
    elif report_type == "driver_details":
        rows = await _driver_detail_rows(
            start_date=start_date, end_date=end_date, driver_oid=driver_oid,
            driver_account_status=driver_account_status, device_enrollment_status=device_enrollment_status,
        )
        _create_detail_sheet(
            workbook, name="Driver Details", title="Driver Details Report", filters=filters,
            headers=DRIVER_DETAILS_HEADERS, rows=rows,
        )
    else:
        rows = await _route_details_report_rows(
            start_date=start_date, end_date=end_date, vehicle_oid=vehicle_oid, driver_oid=driver_oid,
            route_status=route_status, delivery_period=delivery_period, consolidation_id=consolidation_id,
        )
        _create_detail_sheet(
            workbook, name="Route Details", title="Route Details Report", filters=filters,
            headers=ROUTE_DETAILS_HEADERS, rows=rows,
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


async def build_xlsx_report(start_date: datetime | None = None, end_date: datetime | None = None) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise HTTPException(status_code=503, detail="XLSX report dependency is unavailable") from exc
    data = await collect_report_data(start_date, end_date)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["FleetOpt AI Operational Report", datetime.now(timezone.utc).isoformat()])
    summary.append(["Section", "Records"])
    for name, rows in data.items():
        summary.append([name, len(rows)])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B8F")
    for name, rows in data.items():
        sheet = workbook.create_sheet(name)
        headers = list(SECTIONS[name][1])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173B8F")
        for row in rows:
            sheet.append([_display(row.get(header)) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


async def build_pdf_report(start_date: datetime | None = None, end_date: datetime | None = None) -> bytes:
    data = await collect_report_data(start_date, end_date)
    lines = [
        "FleetOpt AI Operational Report",
        f"Generated: {datetime.now(timezone.utc).isoformat()}",
        "",
        "Operational summary",
        *[f"{name}: {len(rows)} record(s)" for name, rows in data.items()],
        "",
        "Model limitations",
        "RUL estimates are planning signals, not exact service dates.",
        "Graph-enhanced demand output requires comparison with the baseline model.",
    ]
    for name, rows in data.items():
        lines.extend(["", name, " | ".join(SECTIONS[name][1])])
        for row in rows[:80]:
            lines.append(" | ".join(str(_display(row.get(field)))[:40] for field in SECTIONS[name][1]))
    return build_text_pdf(lines)


def build_text_pdf(lines: list[str]) -> bytes:
    """Build a standards-compliant, dependency-free A4 PDF with paginated text."""
    def escape(value: str) -> str:
        return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("latin-1", "replace").decode("latin-1")

    pages = [lines[index:index + 48] for index in range(0, len(lines), 48)] or [[]]
    objects: list[bytes] = []
    page_ids: list[int] = []
    font_id = 3
    next_id = 4
    content_objects: list[tuple[int, bytes]] = []
    for page in pages:
        page_id, content_id = next_id, next_id + 1
        next_id += 2
        page_ids.append(page_id)
        commands = ["BT", "/F1 10 Tf", "50 790 Td", "13 TL"]
        for line in page:
            commands.append(f"({escape(str(line))}) Tj")
            commands.append("T*")
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1")
        content_objects.append((content_id, b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"))
        objects.append(f"{page_id} 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>\nendobj\n".encode())
    base_objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        f"2 0 obj\n<< /Type /Pages /Kids [{' '.join(f'{page_id} 0 R' for page_id in page_ids)}] /Count {len(page_ids)} >>\nendobj\n".encode(),
        b"3 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
    ]
    by_id: dict[int, bytes] = {1: base_objects[0], 2: base_objects[1], 3: base_objects[2]}
    for raw in objects:
        object_id = int(raw.split(b" ", 1)[0])
        by_id[object_id] = raw
    for object_id, body in content_objects:
        by_id[object_id] = f"{object_id} 0 obj\n".encode() + body + b"\nendobj\n"
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_id in range(1, max(by_id) + 1):
        offsets.append(len(output))
        output.extend(by_id[object_id])
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    output.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(output)
