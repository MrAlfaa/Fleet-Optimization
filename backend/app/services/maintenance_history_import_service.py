from __future__ import annotations

from datetime import date, datetime, time, timezone
from io import BytesIO
from typing import Any
from uuid import uuid4

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from app.models.maintenance_model import MAINTENANCE_EVENT_KINDS, MAINTENANCE_SYSTEMS
from app.core.database import get_database
from app.models.maintenance_model import MAINTENANCE_COLLECTION, build_maintenance_document
from app.models.user_model import utc_now
from app.models.vehicle_model import VEHICLE_COLLECTION
from app.utils.object_id import validate_object_id


MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
HISTORY_TEMPLATE_HEADERS = (
    "registration_number",
    "service_date",
    "mileage_at_service",
    "maintenance_system",
    "event_kind",
    "work_performed_or_issue",
    "technician_name",
    "service_center",
    "actual_cost_lkr",
    "notes",
)


def build_maintenance_history_template(registration_number: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Maintenance History"
    instructions = workbook.create_sheet("Instructions")
    allowed = workbook.create_sheet("Allowed Values")

    sheet.append(HISTORY_TEMPLATE_HEADERS)
    sheet.append([registration_number.upper(), None, None, None, None, None, None, None, None, None])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{MAX_IMPORT_ROWS + 2}"
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    widths = [22, 16, 20, 24, 22, 42, 24, 26, 18, 42]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    instructions["A1"] = "FleetOpt Maintenance History Import"
    instructions["A1"].font = Font(size=16, bold=True, color="1D4ED8")
    guidance = [
        "Use one row per completed historical service, repair, breakdown, or inspection.",
        "Do not change sheet names or column headers.",
        "service_date must be YYYY-MM-DD and cannot be in the future.",
        "mileage_at_service cannot exceed the vehicle's current mileage.",
        "Use only the dropdown values supplied for maintenance_system and event_kind.",
        "The import is atomic: if any row is invalid, no records are imported.",
        "Delete the blank starter row if it is not needed; blank rows are ignored.",
    ]
    for row, value in enumerate(guidance, start=3):
        instructions.cell(row=row, column=1, value=value)
    instructions.column_dimensions["A"].width = 110

    systems = sorted(MAINTENANCE_SYSTEMS)
    kinds = sorted(MAINTENANCE_EVENT_KINDS)
    allowed.append(["maintenance_system", "event_kind"])
    for index in range(max(len(systems), len(kinds))):
        allowed.append([systems[index] if index < len(systems) else None, kinds[index] if index < len(kinds) else None])
    allowed.sheet_state = "hidden"
    system_range = f"{quote_sheetname(allowed.title)}!$A$2:$A${len(systems) + 1}"
    kind_range = f"{quote_sheetname(allowed.title)}!$B$2:$B${len(kinds) + 1}"
    system_validation = DataValidation(type="list", formula1=system_range, allow_blank=False)
    kind_validation = DataValidation(type="list", formula1=kind_range, allow_blank=False)
    sheet.add_data_validation(system_validation)
    sheet.add_data_validation(kind_validation)
    system_validation.add(f"D2:D{MAX_IMPORT_ROWS + 2}")
    kind_validation.add(f"E2:E{MAX_IMPORT_ROWS + 2}")

    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _text(value: Any) -> str:
    return str(value or "").strip()


def _date(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    else:
        try:
            parsed = datetime.strptime(_text(value), "%Y-%m-%d")
        except ValueError:
            return None
    return parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed.astimezone(timezone.utc)


def parse_maintenance_history_workbook(
    content: bytes, registration_number: str, *, current_mileage: float
) -> list[dict[str, Any]]:
    if not content or len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Workbook must be 5 MB or smaller")
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=False)
        sheet = workbook["Maintenance History"]
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Invalid FleetOpt XLSX workbook") from exc
    headers = tuple(_text(cell.value) for cell in sheet[1])
    if headers != HISTORY_TEMPLATE_HEADERS:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Workbook headers do not match the FleetOpt template")

    parsed_rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    expected_registration = registration_number.strip().upper()
    today = datetime.now(timezone.utc).date()
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values[1:]):
            continue
        if len(parsed_rows) >= MAX_IMPORT_ROWS:
            errors.append({"row": row_number, "errors": [f"A maximum of {MAX_IMPORT_ROWS} rows is allowed"]})
            break
        registration, service_raw, mileage_raw, system_raw, kind_raw, issue_raw, technician, center, cost_raw, notes = values[:10]
        row_errors: list[str] = []
        if _text(registration).upper() != expected_registration:
            row_errors.append("registration_number must match the selected vehicle")
        service_date = _date(service_raw)
        if service_date is None:
            row_errors.append("service_date must use YYYY-MM-DD")
        elif service_date.date() > today:
            row_errors.append("service_date cannot be in the future")
        try:
            mileage = float(mileage_raw)
            if mileage < 0 or mileage > current_mileage:
                row_errors.append("mileage_at_service must be between 0 and current vehicle mileage")
        except (TypeError, ValueError):
            mileage = 0.0
            row_errors.append("mileage_at_service must be a number")
        maintenance_system = _text(system_raw).lower()
        if maintenance_system not in MAINTENANCE_SYSTEMS:
            row_errors.append("maintenance_system is not an allowed value")
        event_kind = _text(kind_raw).lower()
        if event_kind not in MAINTENANCE_EVENT_KINDS:
            row_errors.append("event_kind is not an allowed value")
        issue = _text(issue_raw)
        if not issue:
            row_errors.append("work_performed_or_issue is required")
        try:
            actual_cost = None if cost_raw in (None, "") else float(cost_raw)
            if actual_cost is not None and actual_cost < 0:
                row_errors.append("actual_cost_lkr cannot be negative")
        except (TypeError, ValueError):
            actual_cost = None
            row_errors.append("actual_cost_lkr must be a number")
        for value, label in ((issue_raw, "work_performed_or_issue"), (technician, "technician_name"), (center, "service_center"), (notes, "notes")):
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                row_errors.append(f"{label} cannot start with a spreadsheet formula character")
        if row_errors:
            errors.append({"row": row_number, "errors": row_errors})
            continue
        maintenance_type = {
            "scheduled_service": "preventive",
            "corrective_repair": "corrective",
            "breakdown": "corrective",
            "inspection": "inspection",
        }[event_kind]
        parsed_rows.append({
            "part_name": maintenance_system.replace("_", " ").title(),
            "maintenance_type": maintenance_type,
            "maintenance_system": maintenance_system,
            "event_kind": event_kind,
            "issue_type": issue,
            "service_date": service_date,
            "mileage_at_service": mileage,
            "technician_name": _text(technician) or None,
            "service_center": _text(center) or None,
            "cost": actual_cost,
            "actual_cost": actual_cost,
            "status": "completed",
            "work_order_status": "completed",
            "priority": "medium",
            "notes": _text(notes) or None,
            "completion_evidence": "Imported completed maintenance history",
            "completed_at": service_date,
            "record_source": "imported_history",
        })
    if errors:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail={"message": "Maintenance history workbook contains invalid rows", "errors": errors[:100]},
        )
    if not parsed_rows:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Workbook contains no maintenance history rows")
    return parsed_rows


def _database():
    database = get_database()
    if database is None:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="MongoDB is not configured or connected")
    return database


async def get_vehicle_history_template(vehicle_id: str) -> tuple[bytes, str]:
    try:
        vehicle_oid = validate_object_id(vehicle_id)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid vehicle ID") from exc
    vehicle = await _database()[VEHICLE_COLLECTION].find_one({"_id": vehicle_oid})
    if vehicle is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Vehicle not found")
    registration = str(vehicle.get("registration_number") or vehicle.get("vehicle_code") or "vehicle").upper()
    filename_registration = "".join(character if character.isalnum() else "-" for character in registration).strip("-")
    return build_maintenance_history_template(registration), f"{filename_registration}-maintenance-history-template.xlsx"


async def import_vehicle_maintenance_history(
    vehicle_id: str,
    filename: str | None,
    content: bytes,
    current_user: dict[str, Any],
) -> dict[str, Any]:
    if not filename or not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Upload the FleetOpt .xlsx template")
    try:
        vehicle_oid = validate_object_id(vehicle_id)
        actor_id = validate_object_id(str(current_user.get("id") or current_user.get("_id") or ""))
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid vehicle or user ID") from exc
    database = _database()
    vehicle = await database[VEHICLE_COLLECTION].find_one({"_id": vehicle_oid})
    if vehicle is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Vehicle not found")
    registration = str(vehicle.get("registration_number") or "").upper()
    rows = parse_maintenance_history_workbook(
        content,
        registration,
        current_mileage=float(vehicle.get("current_mileage") or 0),
    )

    duplicate_rows: list[int] = []
    for index, row in enumerate(rows, start=1):
        duplicate = await database[MAINTENANCE_COLLECTION].find_one({
            "vehicle_id": vehicle_oid,
            "status": "completed",
            "service_date": row["service_date"],
            "mileage_at_service": row["mileage_at_service"],
            "maintenance_system": row["maintenance_system"],
            "event_kind": row["event_kind"],
            "issue_type": row["issue_type"],
        }, {"_id": 1})
        if duplicate is not None:
            duplicate_rows.append(index + 1)
    if duplicate_rows:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"History contains records already stored for template row(s): {', '.join(map(str, duplicate_rows))}",
        )

    batch_id = str(uuid4())
    documents = []
    for row in rows:
        prepared = {
            **row,
            "vehicle_id": vehicle_oid,
            "import_batch_id": batch_id,
            "completed_by": actor_id,
        }
        documents.append(build_maintenance_document(prepared, actor_id))

    try:
        result = await database[MAINTENANCE_COLLECTION].insert_many(documents, ordered=True)
        latest_service_date = max(row["service_date"] for row in rows)
        update = await database[VEHICLE_COLLECTION].update_one(
            {"_id": vehicle_oid},
            {
                "$max": {"last_service_date": latest_service_date},
                "$set": {"updated_at": utc_now()},
            },
        )
        if update.matched_count == 0:
            raise RuntimeError("Vehicle changed during maintenance history import")
    except Exception as exc:
        await database[MAINTENANCE_COLLECTION].delete_many({"import_batch_id": batch_id})
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="Maintenance history import failed; no records were retained") from exc

    return {
        "vehicle_id": vehicle_id,
        "registration_number": registration,
        "imported_count": len(result.inserted_ids),
        "import_batch_id": batch_id,
        "latest_service_date": latest_service_date,
    }
